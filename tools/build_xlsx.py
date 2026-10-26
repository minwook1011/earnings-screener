# -*- coding: utf-8 -*-
"""회사별 표준 재무제표 엑셀(docs/data/jp/xlsx/<코드>.xlsx) — 토큰 없이 파이썬만으로.

  자료  결산단신 XBRL(TDnet → 거래소 보관소 www2.jpx.co.jp/disc). 목록은 yanoshin TDnet API.
        첨부 재무제표(손익·포괄손익·재무상태표·현금흐름표·부문)의 **항목 전부**를 싣는다.
  양식  테크윙 재무모델 양식 그대로(tools/fmodel.py = 사용자 「재무모델링」 스킬 build_model.py 사본, 고정).
        '연간' 시트: 12분기 + 10개년 + 5개년 추정(MW) 칸 · 목표주가 · 자산가치 · 기업가치.
        '재무상태표' 시트: 5분기 + 5개년 · 현금흐름 요약 · 재무 비율.
        두 시트 아래에 FPS 식 상세 표(같은 열): 손익·포괄손익·부문·현금흐름 전 항목 / 재무상태표 전 항목.
        단위 백만엔(공시 단위).
  한국어 항목명은 tools/jp_labels_ko.json(일본어 → 한국어). 없는 말은 tools/jp_labels_missing.json 에 모이고
        엑셀에는 영어 요소명이 임시로 들어간다 → 사전에 채운 뒤 다시 돌리면 된다.
  섹터 자료  리서치 폴더의 <코드>_data.json(2단계 보고서 만들 때 모은 월차·수주·수출입 등)과 사이트 월차를
        '섹터 데이터' 시트에 표로 싣고, 출처는 맨 오른쪽 칸.

  python tools/build_xlsx.py 6954 [6999 …]      만들기(결산단신 zip 은 tools/xbrl_cache 에 보관)
"""
import glob
import json
import os
import re
import sys
import time
import urllib.request
import zipfile
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fmodel  # noqa: E402  테크윙 양식(재무모델링 스킬 build_model.py 사본)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, "tools", "xbrl_cache")
LABELS = os.path.join(ROOT, "tools", "jp_labels_ko.json")
MISSING = os.path.join(ROOT, "tools", "jp_labels_missing.json")
OUT = os.path.join(ROOT, "docs", "data", "jp", "xlsx")
PIPE = os.path.join(ROOT, "docs", "data", "jp", "pipeline.json")
MONTHLY = os.path.join(ROOT, "docs", "data", "jp", "monthly")
RESEARCH = os.path.join(os.path.dirname(ROOT), "리서치", "2단계")
KST = timezone(timedelta(hours=9))
LIST = "https://webapi.yanoshin.jp/webapi/tdnet/list/{code}.json?limit=300"
ZIPS = ["https://www.release.tdnet.info/inbs/0812{rest}.zip", "https://www2.jpx.co.jp/disc/{code}0/0812{rest}.zip"]
PDF = "https://www2.jpx.co.jp/disc/{code}0/{doc}.pdf"
NQ = 20      # 분기 칸 최대
NF = 30      # 결산단신 최대 수(약 7년)
STMTS = [("PL", "손익계산서"), ("CI", "포괄손익"), ("BS", "재무상태표"), ("CF", "현금흐름표"), ("SG", "부문")]
KIND = {"bs": "BS", "fs": "BS", "pl": "PL", "pc": "PL", "ci": "CI", "cf": "CF", "sg": "SG"}


# ---------- 받기 ----------
def get(url, tries=3):
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            return urllib.request.urlopen(req, timeout=30).read()
        except Exception:
            if i == tries - 1:
                return None
            time.sleep(2)


_PAGES = {}


def official(code, days=2):
    """TDnet 공식 당일 목록(오늘·어제) — 위 LIST 는 한 시간 넘게 늦어 발표 직후 결산단신이 빠진다(2026-10-26 신에츠폴리머)."""
    out, today = [], datetime.now(KST).date()
    for k in range(days):
        ymd = (today - timedelta(days=k)).strftime("%Y%m%d")
        for page in range(1, 40):
            u = f"https://www.release.tdnet.info/inbs/I_list_{page:03d}_{ymd}.html"
            if u not in _PAGES:   # 한 번 실행 안에서 여러 회사가 같은 쪽을 다시 받지 않게
                _PAGES[u] = get(u, tries=1)
            raw = _PAGES[u]
            if not raw:
                break
            s = raw.decode("utf-8", "ignore")
            rows = re.findall(r"<tr>(.*?)</tr>", s, re.S)
            if not any("kjCode" in r for r in rows):
                break
            for r in rows:
                if f">{code}0</td>" not in r or "kjXbrl" not in r:
                    continue
                t = re.search(r'kjTitle[^>]*><a href="(\d{18})\.pdf"[^>]*>([^<]*)</a>', r)
                if t and "決算短信" in t.group(2) and "訂正" not in t.group(2) and ".zip" in r:
                    out.append({"doc": t.group(1), "date": f"{ymd[:4]}-{ymd[4:6]}-{ymd[6:]}", "title": t.group(2).strip()})
    return out


def filings(code):
    raw = get(LIST.format(code=code))
    if not raw:
        sys.exit("결산단신 목록을 못 받았습니다")
    out = []
    for it in json.loads(raw).get("items", []):
        t = it.get("Tdnet", it) if isinstance(it, dict) else {}
        title = t.get("title", "")
        if "決算短信" not in title or "訂正" in title or not t.get("url_xbrl"):
            continue
        m = re.search(r"(\d{18})\.pdf", t.get("document_url", ""))
        if m:
            out.append({"doc": m.group(1), "date": t.get("pubdate", "")[:10], "title": title})
    have = {f["doc"] for f in out}
    out = sorted([f for f in official(code) if f["doc"] not in have] + out, key=lambda f: (f["date"], f["doc"]), reverse=True)
    keep, n_year = [], 0
    for f in out:                     # 새것부터
        annual = not re.search(r"四半期|中間", f["title"])
        if len(keep) < 16 or (annual and n_year < 9):
            keep.append(f)
            n_year += annual
    return keep


def zip_of(code, doc):
    p = os.path.join(CACHE, code, doc + ".zip")
    if os.path.exists(p):
        if zipfile.is_zipfile(p):
            return p
        os.remove(p)   # 받다 끊긴 파일
    for u in ZIPS:
        b = get(u.format(code=code, rest=doc[4:]))
        if b and b[:2] == b"PK":
            os.makedirs(os.path.dirname(p), exist_ok=True)
            open(p + ".part", "wb").write(b)
            os.replace(p + ".part", p)
            return p
    return None


# ---------- 읽기 ----------
class IX(HTMLParser):
    """인라인 XBRL 한 장 — 문맥(기간·멤버)과 표의 행(이름표 + 숫자들)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.ctx, self.rows = {}, []
        self.cur = self.field = self.row = self.cell = self.fact = None
        self.buf = ""

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "xbrli:context":
            self.cur = {"start": None, "end": None, "inst": None, "mem": [], "id": a.get("id")}
        elif self.cur is not None and tag in ("xbrli:startdate", "xbrli:enddate", "xbrli:instant", "xbrldi:explicitmember"):
            self.field, self.buf = tag, ""
        elif tag == "tr":
            self.row = {"cells": [], "facts": []}
        elif tag in ("td", "th") and self.row is not None:
            self.cell = ""
        elif tag == "ix:nonfraction":
            self.fact = {"name": a.get("name", ""), "ctx": a.get("contextref", ""), "scale": int(a.get("scale") or 0),
                         "sign": a.get("sign"), "fmt": a.get("format", ""), "txt": ""}

    def handle_endtag(self, tag):
        if self.field == tag:
            k = {"xbrli:startdate": "start", "xbrli:enddate": "end", "xbrli:instant": "inst"}.get(tag)
            if k:
                self.cur[k] = self.buf.strip()
            else:
                self.cur["mem"].append(self.buf.strip())
            self.field = None
        elif tag == "xbrli:context" and self.cur:
            self.ctx[self.cur["id"]] = self.cur
            self.cur = None
        elif tag == "ix:nonfraction" and self.fact is not None:
            f = self.fact
            t = f["txt"].strip()
            if "zero" in f["fmt"] or t in ("", "－", "-", "―", "—", "‐"):
                v = 0.0
            else:
                try:
                    v = float(re.sub(r"[^\d.]", "", t)) * 10 ** f["scale"]
                except ValueError:
                    v = None
            if v is not None and f["sign"] == "-":
                v = -v
            f["v"] = v
            if self.row is not None:
                self.row["facts"].append(f)
            self.fact = None
        elif tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row["cells"].append(self.cell)
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.row["facts"]:
                self.rows.append(self.row)
            self.row = None

    def handle_data(self, d):
        if self.field:
            self.buf += d
        if self.fact is not None:
            self.fact["txt"] += d
        elif self.cell is not None:
            self.cell += d


def clean(s):
    return re.sub(r"\s+", "", (s or "").replace("　", " ")).strip()


def labels_of(z):
    """회사 고유 요소·부문 멤버의 일본어 이름(lab.xml)."""
    out = {}
    for n in z.namelist():
        if n.endswith("-lab.xml"):
            s = z.read(n).decode("utf-8", "ignore")
            locs = dict((m[1], m[0].split("#")[-1]) for m in re.findall(r'<link:loc[^>]*xlink:href="([^"]+)"[^>]*xlink:label="([^"]+)"', s))
            arcs = dict(re.findall(r'<link:labelArc[^>]*xlink:from="([^"]+)"[^>]*xlink:to="([^"]+)"', s))
            labs = {}
            for lab, txt in re.findall(r'<link:label[^>]*xlink:label="([^"]+)"[^>]*xml:lang="ja"[^>]*>([^<]*)<', s):
                labs.setdefault(lab, txt)
            for frm, to in arcs.items():
                if frm in locs and to in labs:
                    out[locs[frm]] = labs[to]
                    out.setdefault(locs[frm].split("_")[-1], labs[to])
    return out


def months(a, b):
    return (b.year - a.year) * 12 + b.month - a.month + (1 if b.day >= 28 else 0)


def plabel(end, fsm):
    off = (end.month - fsm) % 12
    q = off // 3 + 1
    fy_start = end.year if end.month >= fsm else end.year - 1
    fy_end = fy_start + (1 if fsm != 1 else 0)
    return fy_end, q


STD = ("jppfs_cor", "jpigp_cor", "jpcrp_cor", "jpdei_cor", "tse-ed-t")


def norm(name):
    pre, _, loc = name.rpartition(":")
    return name if pre in STD else "co:" + loc


def parse_filing(path):
    z = zipfile.ZipFile(path)
    jlab = labels_of(z)
    out, parts, ctx = [], [], {}
    for n in z.namelist():    # 문맥(기간·멤버)은 한 파일에만 정의돼 있기도 해서 먼저 다 모은다
        m = re.search(r"/\d{7}-[a-z]{2}(bs|fs|pl|pc|ci|cf|sg)\d+-.*-ixbrl\.htm$", n)
        if not m:
            continue
        p = IX()
        p.feed(z.read(n).decode("utf-8", "ignore"))
        ctx.update(p.ctx)
        parts.append((m, p))
    fsm = next((int(c["start"][5:7]) for c in ctx.values() if c["start"] and c["end"]
                and months(date.fromisoformat(c["start"]), date.fromisoformat(c["end"])) >= 3), None)
    for m, p in parts:
        for order, row in enumerate(p.rows):
            lab = next((clean(x) for x in row["cells"] if clean(x) and not re.fullmatch(r"[\d,.\-△▲()（）%－―]+", clean(x))), "")
            for f in row["facts"]:
                c = ctx.get(f["ctx"])
                if not c or f["v"] is None:
                    continue
                out.append({"stmt": KIND[m.group(1)], "name": norm(f["name"]), "lab": lab, "order": order, "c": c,
                            "v": f["v"], "fsm": fsm,
                            "mem": [norm(x) for x in c["mem"] if not x.endswith("NonConsolidatedMember")]})   # 개별만 내는 회사
    return out, jlab


STDMEM = {"ReportableSegmentsMember": "보고부문 합계", "ReconcilingItemsMember": "조정액",
          "TotalOfReportableSegmentsAndOthersMember": "합계",
          "OperatingSegmentsNotIncludedInReportableSegmentsAndOtherRevenueGeneratingBusinessActivitiesMember": "기타"}


# ---------- 모으기 ----------
def collect(code):
    fl = filings(code)
    data = {k: {"rows": [], "seen": {}, "ytd": {}, "inst": {}} for k, _ in STMTS}
    jlab_all, used, fsm0 = {}, [], None
    for f in fl:                                  # 새것부터 — 먼저 들어간 값이 우선
        path = zip_of(code, f["doc"])
        if not path:
            continue
        facts, jlab = parse_filing(path)
        if not facts:
            continue
        used.append(f)
        jlab_all.update(jlab)
        prev = {k: None for k, _ in STMTS}
        for x in facts:
            st, c = x["stmt"], x["c"]
            if st != "SG" and x["mem"]:
                continue                          # 단체·기타 축은 뺌(연결만)
            key = x["name"] + ("|" + "|".join(x["mem"]) if x["mem"] else "")
            D = data[st]
            if key not in D["seen"]:
                lab = x["lab"]
                if x["mem"]:
                    mem = x["mem"][-1].split(":")[-1]
                    lab = (STDMEM.get(mem) or jlab.get(mem) or jlab_all.get(mem) or mem) + " · " + x["lab"]
                pos = D["rows"].index(prev[st]) + 1 if prev[st] in D["rows"] else len(D["rows"])
                D["rows"].insert(pos, key)
                D["seen"][key] = {"lab": lab, "name": x["name"], "mem": x["mem"]}
            prev[st] = key
            if c["inst"]:
                if not x["fsm"] and not fsm0:
                    continue
                e = date.fromisoformat(c["inst"])
                D["inst"].setdefault(key, {}).setdefault(plabel(e, x["fsm"] or fsm0), x["v"])
            elif c["start"] and c["end"]:
                s, e = date.fromisoformat(c["start"]), date.fromisoformat(c["end"])
                fsm0 = fsm0 or s.month
                q = round(months(s, e) / 3)
                if 1 <= q <= 4:
                    fy, _ = plabel(e, s.month)
                    D["ytd"].setdefault(key, {}).setdefault((fy, q), x["v"])
    return data, used, fsm0


def quarters(data):
    ks = set()
    for st in ("PL", "BS", "CF", "SG", "CI"):
        for d in list(data[st]["ytd"].values()) + list(data[st]["inst"].values()):
            ks |= set(d)
    ks = sorted(ks)
    return ks[-NQ:]


def standalone(ytd, fy, q):
    a = ytd.get((fy, q))
    if a is None:
        return None
    if q == 1:
        return a
    b = ytd.get((fy, q - 1))
    return None if b is None else a - b


# ---------- 한국어 ----------
def ko_dict():
    try:
        return json.load(open(LABELS, encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def eng(name):
    n = name.split(":")[-1]
    n = re.sub(r"(IFRS|Summary|CY|NC)$", "", n)
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", n)


def ko(jp, name, KO, miss):
    parts = [p for p in jp.split(" · ")]
    out = []
    for i, p in enumerate(parts):
        p = re.sub(r"、?報告セグメント|\s*\[メンバー\]|メンバー$", "", p).strip()
        if not p:
            continue
        alt = [p] + [re.sub(r"四半期|中間|当期", w, p) for w in ("中間", "当期", "")] + [re.sub(r"（[^）]*△[^）]*）$", "", p)]
        hit = next((KO[a] for a in alt if a in KO), None)
        if hit:
            out.append(hit)
        elif re.search(r"[぀-ヿ一-鿿]", p):
            miss.add(p)
            out.append(eng(re.sub(r"(Reportable)?(Segment)?Member$", "", name)) if i == len(parts) - 1
                       else p.encode("ascii", "ignore").decode() or "부문")
        else:
            out.append(p)
    return " · ".join(out) or eng(name)


# ---------- 엑셀 공통(섹터 데이터 시트) ----------
F = "맑은 고딕"
GRAY = PatternFill("solid", fgColor="FFD9D9D9")
BAND = PatternFill("solid", fgColor="FF808080")
THIN = Side(style="thin", color="FFBFBFBF")
NUM = '#,##0.0;\\(#,##0.0\\);"-"'
PCT = '0.0%;\\-0.0%;"-"'
CEN = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")



def band(ws, r, text, ncol):
    for i in range(1, ncol + 1):
        c = ws.cell(r, i, text if i == 1 else None)
        c.font = Font(name=F, size=11, bold=True, color="FFFFFFFF")
        c.fill = BAND
        c.alignment = CEN



def table_block(ws, r, title, header, rows, src, width):
    """섹터 자료 표 하나 — 회색 띠 제목, 머리줄, 행들, 맨 오른쪽 출처."""
    band(ws, r, title, width)
    r += 1
    for i, h in enumerate(["항목"] + header, 1):
        c = ws.cell(r, i, h)
        c.font = Font(name=F, size=10, bold=True)
        c.fill = GRAY
        c.alignment = CEN
    c = ws.cell(r, width, "출처")
    c.font = Font(name=F, size=10, bold=True)
    c.fill = GRAY
    c.alignment = CEN
    r += 1
    first = True
    for lab, vals in rows:
        ws.cell(r, 1, kt(lab)).font = Font(name=F, size=11)
        for i, v in enumerate(vals):
            c = ws.cell(r, 2 + i, v)
            c.font = Font(name=F, size=11)
            c.alignment = CEN
            if isinstance(v, (int, float)):
                c.number_format = '#,##0.0;\\(#,##0.0\\);"-"' if isinstance(v, float) else "#,##0"
        if first:
            s = ws.cell(r, width, src or "2단계 보고서 자료(회사 설명회 자료·공시)")
            s.font = Font(name=F, size=9, color="FF808080")
            first = False
        r += 1
    return r + 1


KEYW = {"rev": "매출", "revenue": "매출", "sales": "매출", "op": "영업이익", "ni": "순이익", "np": "순이익", "eps": "EPS",
        "n": "애널리스트 수", "n_eps": "EPS 애널리스트 수", "n_rev": "매출 애널리스트 수", "eps_90d_ago": "90일 전 EPS",
        "mean": "평균", "median": "중앙값", "high": "최고", "low": "최저", "target": "목표주가", "consensus": "컨센서스",
        "guidance": "가이던스", "prev": "직전", "h1": "상반기", "h2": "하반기", "orders": "수주", "order": "수주",
        "backlog": "수주잔고", "segments": "부문별", "segment": "부문", "region": "지역별", "share": "비중", "yoy": "YoY",
        "qoq": "QoQ", "bridge": "증감 요인", "total": "합계", "monthly": "월차", "export": "수출", "exports": "수출",
        "import": "수입", "imports": "수입", "trade": "무역", "machine": "공작기계", "tool": "", "price": "가격",
        "margin": "이익률", "opm": "영업이익률", "gpm": "매출총이익률", "units": "수량", "volume": "물량", "cagr": "CAGR",
        "fy": "FY", "actual": "실적", "plan": "계획", "forecast": "예상", "est": "추정", "peers": "비교 기업",
        "valuation": "밸류에이션", "per": "PER", "pbr": "PBR", "roe": "ROE", "mc": "시가총액", "capex": "설비투자",
        "rd": "연구개발비", "da": "감가상각비", "fx": "환율", "usd": "달러", "eur": "유로", "cny": "위안",
        "q1": "1Q", "q2": "2Q", "q3": "3Q", "q4": "4Q", "kpi": "핵심 지표", "stores": "점포", "store": "점포",
        "customers": "고객 수", "ticket": "객단가", "existing": "기존점", "all": "전체", "growth": "성장률", "pct": "%"}


def kt(k):
    """영어 키 → 한국어(모르는 말은 그대로)."""
    k = str(k)
    if k.lower() in KEYW:
        return KEYW[k.lower()]
    if re.fullmatch(r"[A-Za-z0-9_ ·\-]+", k) and ("_" in k or " · " in k):
        out = []
        for part in k.split(" · "):
            toks = [KEYW.get(t.lower(), t) for t in part.split("_")]
            out.append(" ".join(t for t in toks if t))
        return " · ".join(out)
    return k


def is_num_list(v):
    return isinstance(v, list) and v and all(isinstance(x, (int, float)) or x is None for x in v)


def sector_tables(code):
    """리서치 폴더 <코드>_data.json 에서 표 모양 자료를 뽑는다(재무 숫자는 재무제표 시트가 있으니 뺌)."""
    out = []
    files = glob.glob(os.path.join(RESEARCH, code + " *", code + "_data.json"))
    try:
        d = json.load(open(files[0], encoding="utf-8")) if files else None
    except ValueError:
        d = None   # 깨진 자료 파일은 건너뜀
    if isinstance(d, dict):
        skip = {"quarterly", "annual", "price", "company", "code", "fiscal", "unit", "asof"}
        for k, v in d.items():
            if k in skip or not isinstance(v, dict):
                continue
            src = next((str(v[s]) for s in ("src", "source", "출처", "sources") if s in v), "")
            note = str(v.get("note", "")) if isinstance(v.get("note"), str) else ""
            src = " · ".join(x for x in (src, note) if x)
            if isinstance(v.get("labels"), list):
                rows = [(kk, vv) for kk, vv in v.items() if is_num_list(vv)]
                if rows:
                    out.append((k, [str(x) for x in v["labels"]], rows, src))
                continue
            flat = [(kk, vv) for kk, vv in v.items() if isinstance(vv, (int, float))]
            if len(flat) >= 2:
                out.append((k, ["값"], [(kk, [vv]) for kk, vv in flat], src))
            for kk, vv in v.items():
                if isinstance(vv, dict):
                    f2 = [(a, b) for a, b in vv.items() if isinstance(b, (int, float))]
                    if len(f2) >= 2:
                        out.append((k + " · " + kk, ["값"], [(a, [b]) for a, b in f2], src))
        for k, v in d.items():   # 표 모양 리스트([{…}, …])
            if k in skip or not isinstance(v, list) or not v or not all(isinstance(x, dict) for x in v):
                continue
            cols = [c for c in v[0].keys()]
            if len(cols) < 2:
                continue
            out.append((k, cols[1:], [(x.get(cols[0]), [x.get(c) for c in cols[1:]]) for x in v], ""))
    mp = os.path.join(MONTHLY, code + ".json")
    if os.path.exists(mp):
        m = json.load(open(mp, encoding="utf-8"))
        rows = m.get("rows", [])[-24:]
        out.insert(0, ("월차", [r[0] for r in rows], [(c, [r[i + 1] for r in rows]) for i, c in enumerate(m.get("cols", []))],
                       (m.get("src") or "회사 월차 공시") + (" · " + m["note"] if m.get("note") else "")))
    return out


# ---------- 테크윙 양식(사용자 재무모델링 스킬의 build_model.py = tools/fmodel.py) + 아래 상세 표 ----------
SCALE, UNIT = 1e6, "백만엔"     # 일본 공시 단위 그대로
NQ_T, NY_T, NB_Q, NB_Y = 12, 10, 5, 5   # 양식: 분기 12 · 연간 10 · 재무상태표 분기 5 · 연간 5

# 양식 키 ← XBRL 요소(로컬 이름, 앞에서부터 먼저 찾은 것). 일본 기준·IFRS 둘 다
PLMAP = {
    "revenue": ["NetSales", "Revenue", "RevenueIFRS", "NetSalesIFRS", "OperatingRevenue1", "OperatingRevenue2", "Revenues",
                "OperatingRevenuesSE", "NetSalesOfCompletedConstructionContracts", "OrdinaryIncomeBNK"],
    "cogs": ["CostOfSales", "CostOfSalesIFRS", "OperatingCost", "CostOfSalesOfCompletedConstructionContracts"],
    "gp": ["GrossProfit", "GrossProfitIFRS", "OperatingGrossProfit"],
    "sga": ["SellingGeneralAndAdministrativeExpenses", "SellingGeneralAndAdministrativeExpensesIFRS", "TotalSellingGeneralAndAdministrativeExpenses"],
    "op": ["OperatingIncome", "OperatingProfitLossIFRS", "OperatingIncomeLoss"],
    "oth_inc": ["ExtraordinaryIncome", "OtherIncomeIFRS"],
    "oth_exp": ["ExtraordinaryLoss", "OtherExpensesIFRS"],
    "fin_inc": ["NonOperatingIncome", "FinanceIncomeIFRS"],
    "fin_exp": ["NonOperatingExpenses", "FinanceCostsIFRS"],
    "pbt": ["IncomeBeforeIncomeTaxes", "ProfitLossBeforeTaxIFRS", "IncomeLossBeforeIncomeTaxes"],
    "tax": ["IncomeTaxes", "IncomeTaxExpenseIFRS", "TotalIncomeTaxes"],
    "ni": ["ProfitLoss", "ProfitLossIFRS", "NetIncome", "NetIncomeLoss"],
    "ni_owner": ["ProfitLossAttributableToOwnersOfParent", "ProfitLossAttributableToOwnersOfParentIFRS"],
    "salary": ["SalariesAndAllowancesSGA", "SalariesSGA", "EmployeesSalariesAndBonusesSGA", "SalariesAndWagesSGA",
               "SalaryAndAllowancesSGA", "PersonnelExpensesSGA"],
    "dep": ["DepreciationSGA", "DepreciationAndAmortizationSGA"],
    "fees": ["CommissionFeeSGA", "CommissionFeesSGA", "CommissionSGA"],
    "rnd": ["ResearchAndDevelopmentExpensesSGA", "ResearchAndDevelopmentCostsSGA"],
}
BSMAP = {
    "current_assets": ["CurrentAssets", "CurrentAssetsIFRS"],
    "cash": ["CashAndDeposits", "CashAndCashEquivalentsIFRS", "CashAndDue"],
    "st_fin": ["ShortTermInvestmentSecurities", "SecuritiesCA", "OtherFinancialAssetsCAIFRS"],
    "receivables": ["NotesAndAccountsReceivableTradeAndContractAssets", "NotesAndAccountsReceivableTrade",
                    "AccountsReceivableTradeAndContractAssets", "AccountsReceivableTrade", "TradeAndOtherReceivablesCAIFRS"],
    "inventory": ["Inventories", "InventoriesCAIFRS"],
    "noncurrent_assets": ["NoncurrentAssets", "NonCurrentAssetsIFRS"],
    "ppe": ["PropertyPlantAndEquipment", "PropertyPlantAndEquipmentIFRS"],
    "intangibles": ["IntangibleAssets", "IntangibleAssetsIFRS"],
    "inv_property": ["InvestmentPropertyIFRS", "RealEstateForInvestmentNet", "RealEstateForInvestment"],
    "total_assets": ["Assets", "AssetsIFRS"],
    "current_liab": ["CurrentLiabilities", "TotalCurrentLiabilitiesIFRS", "CurrentLiabilitiesIFRS"],
    "payables": ["NotesAndAccountsPayableTrade", "AccountsPayableTrade", "TradeAndOtherPayablesCLIFRS"],
    "st_debt": ["ShortTermLoansPayable", "ShortTermBorrowings", "BorrowingsCLIFRS", "CommercialPapersLiabilities"],
    "noncurrent_liab": ["NoncurrentLiabilities", "NonCurrentLabilitiesIFRS", "NonCurrentLiabilitiesIFRS"],
    "lt_debt": ["LongTermLoansPayable", "BorrowingsNCLIFRS", "LongTermBorrowings"],
    "bonds": ["BondsPayable", "ConvertibleBondTypeBondsWithSubscriptionRightsToSharesNCL", "BondsAndBorrowingsNCLIFRS"],
    "total_liab": ["Liabilities", "LiabilitiesIFRS"],
    "capital_stock": ["CapitalStock", "ShareCapitalIFRS"],
    "retained": ["RetainedEarnings", "RetainedEarningsIFRS"],
    "nci": ["NonControllingInterests", "NonControllingInterestsIFRS"],
    "total_equity": ["NetAssets", "EquityIFRS"],
}
INV_PARTS = ["MerchandiseAndFinishedGoods", "Merchandise", "FinishedGoods", "WorkInProcess", "RawMaterialsAndSupplies",
             "RawMaterials", "Supplies", "CostsOnUncompletedConstructionContracts", "RealEstateForSale"]
CURRENT_LTD = ["CurrentPortionOfLongTermLoansPayable", "CurrentPortionOfBonds", "CurrentPortionOfConvertibleBonds"]
CFMAP = {
    "ocf": ["NetCashProvidedByUsedInOperatingActivities", "NetCashProvidedByUsedInOperatingActivitiesIFRS"],
    "capex": ["PurchaseOfPropertyPlantAndEquipmentInvCF", "PurchaseOfPropertyPlantAndEquipmentInvCFIFRS"],
    "capex2": ["PurchaseOfIntangibleAssetsInvCF", "PurchaseOfIntangibleAssetsInvCFIFRS", "PurchaseOfSoftwareInvCF"],
    "dividends": ["CashDividendsPaidFinCF", "DividendsPaidFinCFIFRS", "DividendsPaidToOwnersOfParentFinCFIFRS"],
}


def local(name):
    return name.split(":")[-1]


def pick(D, names, src="ytd"):
    """요소 이름 목록(멤버 없는 줄)을 합쳐 {(fy, q): 값}. 요소 이름이 해마다 바뀌는 회사(예 매출 → 영업수익)가 있어
    최근 기간까지 있는 이름을 먼저 쓰고, 빈 기간만 다른 이름으로 메운다."""
    per = []
    for i, n in enumerate(names):
        got = {}
        for k in D["rows"]:
            m = D["seen"][k]
            if not m["mem"] and local(m["name"]) == n:
                for kk, v in D[src].get(k, {}).items():
                    got.setdefault(kk, v)
        if got:
            per.append((max(got), -i, got))
    out = {}
    for _, _, got in sorted(per, reverse=True):
        for kk, v in got.items():
            out.setdefault(kk, v)
    return out


def qseq(last, n):
    fy, q = last
    out = []
    for _ in range(n):
        out.append((fy, q))
        fy, q = (fy, q - 1) if q > 1 else (fy - 1, 4)
    return out[::-1]


def qlab(k):
    return f"{k[0]}.{k[1]}Q"


def shares_of(path):
    """요약 XBRL 의 기말 발행주식수(자기주식 포함)."""
    try:
        z = zipfile.ZipFile(path)
        for n in z.namelist():
            if "Summary" in n and n.endswith("ixbrl.htm"):
                s = z.read(n).decode("utf-8", "ignore")
                m = re.search(r'name="tse-ed-t:NumberOfIssuedAndOutstandingSharesAtTheEndOfFiscalYearIncludingTreasuryStock"[^>]*>([\d,]+)<', s)
                if m:
                    return int(m.group(1).replace(",", ""))
    except (OSError, zipfile.BadZipFile):
        pass
    return None


def price_of(code):
    try:
        r = json.load(open(os.path.join(ROOT, "docs", "data", "jp", "ret.json"), encoding="utf-8"))["items"].get(code)
        if r:
            return r["px"], r["d"].replace("-", ".")
    except (OSError, ValueError, KeyError):
        pass
    return None, None


def make_model(code, name, data, used, fsm, KO, miss):
    PL, BS, CF, SG = data["PL"], data["BS"], data["CF"], data["SG"]
    rev = pick(PL, PLMAP["revenue"])
    if not rev:
        return None
    # 분기·연간 칸
    have = sorted(k for k in rev)
    last = have[-1]
    qk = qseq(last, NQ_T)
    yfy = sorted({fy for fy, q in rev if q == 4})
    yk = list(range(yfy[-1] - NY_T + 1, yfy[-1] + 1)) if yfy else []
    Q = {"labels": [qlab(k) for k in qk]}
    Y = {"labels": yk}
    for key, names in PLMAP.items():
        ytd = pick(PL, names)
        Q[key] = [standalone(ytd, fy, q) for fy, q in qk]
        Y[key] = [ytd.get((fy, 4)) for fy in yk]
    if not any(Y["ni_owner"]) and not any(Q["ni_owner"]):      # 개별만 내는 회사 = 순이익 전부 지배주주
        Q["ni_owner"], Y["ni_owner"] = Q["ni"], Y["ni"]
    for blk in (Q, Y):    # 비용은 양수(환입이면 음수) — XBRL 도 양수로 들어온다
        for k in ("cogs", "sga", "tax"):
            blk[k] = [v for v in blk[k]]
    jgaap = not any(local(PL["seen"][k]["name"]).endswith("IFRS") for k in PL["rows"])
    # 부문(외부 고객 매출, 부문 멤버별)
    segs = []
    seen_mem = []
    for k in SG["rows"]:
        m = SG["seen"][k]
        if not m["mem"] or local(m["name"]) not in ("RevenuesFromExternalCustomers", "RevenueFromExternalCustomersIFRS", "SalesToExternalCustomersIFRS", "NetSales"):
            continue
        mem = local(m["mem"][-1])
        if mem in ("ReportableSegmentsMember", "ReconcilingItemsMember", "TotalOfReportableSegmentsAndOthersMember",
                   "OperatingSegmentsNotIncludedInReportableSegmentsAndOtherRevenueGeneratingBusinessActivitiesMember") and mem != "":
            if mem != "OperatingSegmentsNotIncludedInReportableSegmentsAndOtherRevenueGeneratingBusinessActivitiesMember":
                continue
        if mem in seen_mem:
            continue
        seen_mem.append(mem)
        ytd = {}
        for kk in SG["rows"]:
            mm = SG["seen"][kk]
            if mm["mem"] and local(mm["mem"][-1]) == mem and local(mm["name"]) == local(m["name"]):
                for a, v in SG["ytd"].get(kk, {}).items():
                    ytd.setdefault(a, v)
        jl = m["lab"].split(" · ")[0]
        nm = ko(jl, mem, KO, miss) if re.search(r"[぀-ヿ一-鿿]", jl) else (STDMEM.get(mem) or jl)
        segs.append({"name": nm, "annual": [ytd.get((fy, 4)) for fy in yk],
                     "quarterly": [standalone(ytd, fy, q) for fy, q in qk]})
    # 재무상태표
    def bs_series(keys):
        out = {}
        for key, names in BSMAP.items():
            v = pick(BS, names, "inst")
            out[key] = [v.get(k) for k in keys]
        parts = [pick(BS, [n], "inst") for n in INV_PARTS]
        out["inventory"] = [v if v is not None else (sum(p[k] for p in parts if k in p) if any(k in p for p in parts) else None)
                            for v, k in zip(out["inventory"], keys)]
        cl = [pick(BS, [n], "inst") for n in CURRENT_LTD]
        out["current_ltd"] = [sum(p[k] for p in cl if k in p) if any(k in p for p in cl) else None for k in keys]
        own = pick(BS, ["EquityAttributableToOwnersOfParentIFRS"], "inst")
        sub = pick(BS, ["SubscriptionRightsToShares"], "inst")
        out["equity_owner"] = [own.get(k) if own else (None if te is None else te - (nci or 0) - sub.get(k, 0))
                               for k, te, nci in zip(keys, out["total_equity"], out["nci"])]
        return out
    ta = pick(BS, BSMAP["total_assets"], "inst")
    bq = [k for k in qseq(max(ta), NB_Q)] if ta else []
    by = sorted(fy for fy, q in ta if q == 4)[-NB_Y:] if ta else []
    BQ = {"labels": [qlab(k) for k in bq], **bs_series(bq)}
    BY = {"labels": by, **bs_series([(fy, 4) for fy in by])}
    # 현금흐름(연간 + 분기 TTM)
    cfv = {k: pick(CF, v) for k, v in CFMAP.items()}

    def ttm(ytd, fy, q):
        if q == 4:
            return ytd.get((fy, 4))
        a, b, c = ytd.get((fy, q)), ytd.get((fy - 1, 4)), ytd.get((fy - 1, q))
        return None if None in (a, b, c) else a + b - c

    def cf_block(keys, f):
        ocf = [f(cfv["ocf"], *k) for k in keys]
        c1 = [f(cfv["capex"], *k) for k in keys]
        c2 = [f(cfv["capex2"], *k) for k in keys]
        capex = [None if a is None else abs(a) + abs(b or 0) for a, b in zip(c1, c2)]
        div = [f(cfv["dividends"], *k) for k in keys]
        return {"ocf": ocf, "capex": capex, "dividends": [None if v is None else abs(v) for v in div]}
    CFY = {"labels": by, **cf_block([(fy, 4) for fy in by], lambda d, fy, q: d.get((fy, q)))}
    CFQ = {"labels": BQ["labels"], **cf_block(bq, ttm)}
    px, pdate = price_of(code)
    sh = shares_of(zip_of(code, used[0]["doc"])) if used else None
    end_m = (fsm - 2) % 12 + 1 if fsm else 3
    meta = {
        "company": name, "ticker": code, "market": "JP", "currency": "JPY", "unit_label": UNIT, "scale": SCALE,
        "fiscal_note": f"{end_m}월 결산 · {'일본 기준' if jgaap else 'IFRS'} · {qlab(last)} = {last[0]}년 {end_m}월 결산기 {last[1]}분기",
        "tax_rate": 0.30,
        "sources": ["결산단신 XBRL(TDnet·JPX 보관소) — 쓴 문서는 '출처' 시트", "분기 = 누계에서 앞 분기를 뺀 3개월 값",
                    "주가: 야후 파이낸스 일봉" + (f" {pdate}" if pdate else "")],
    }
    if jgaap:
        meta["labels"] = {"oth_inc": "특별이익", "oth_exp": "특별손실", "fin_inc": "영업외수익", "fin_exp": "영업외비용"}
    model = {"meta": meta, "quarterly": Q, "annual": Y, "segments": segs,
             "forecast": {"mw_years": [yk[-1] + i for i in range(1, 6)] if yk else []},
             "bs": {"quarterly": BQ, "annual": BY}, "cf": {"annual": CFY, "quarterly": CFQ}}
    if px and sh:
        model["market_data"] = {"price": px, "shares": sh, "price_date": pdate}
    return model, qk, yk, bq, by


def detail_block(ws, r, title, cols, last_col, rows, note):
    """FPS 식 상세 표 — 회색 띠 + 열 머리(양식과 같은 칸) + 줄들. rows = [(라벨, {열: 값}, 굵게)]."""
    for c in range(1, last_col + 1):
        cell = ws.cell(r, c, title if c == 1 else None)
        cell.fill, cell.font, cell.alignment = fmodel.fill(fmodel.GRAY_SEC), fmodel.font(True, "FFFFFF"), fmodel.CENTER
    ws.cell(r, 2, note).font = fmodel.font(True, "FFFFFF", 9)
    r += 1
    ws.cell(r, 1, "항목")
    ws.cell(r, 2, "참고")
    for c, kind, l in cols:
        ws.cell(r, c, l)
    for c in range(1, last_col + 1):
        cell = ws.cell(r, c)
        cell.fill, cell.font, cell.alignment = fmodel.fill(fmodel.GRAY_HDR), fmodel.font(True, size=10), fmodel.CENTER
        cell.border = Border(top=fmodel.thin, bottom=fmodel.thin)
    r += 1
    for lab, vals, bold in rows:
        ws.cell(r, 1, lab).font = fmodel.font(bold)
        ws.cell(r, 1).alignment = fmodel.CENTER
        for c, v in vals.items():
            cell = ws.cell(r, c, v)
            cell.number_format = fmodel.NUM
            cell.font = fmodel.font(bold)
            cell.alignment = fmodel.CENTER
        r += 1
    return r + 1


BOLD = re.compile(r"^(NetSales|Revenue|GrossProfit|OperatingIncome|OrdinaryIncome|ProfitLoss|Assets$|Liabilities$|NetAssets|"
                  r"CurrentAssets|NoncurrentAssets|CurrentLiabilities|NoncurrentLiabilities|NetCashProvided|ComprehensiveIncome|"
                  r"IncomeBeforeIncomeTaxes|ShareholdersEquity|LiabilitiesAndNetAssets)")


def stmt_rows(D, KO, miss, colmap, mode, with_mem=False):
    """mode q: 분기 단독, cum: 누계 그대로, inst: 잔액. colmap = {열: (fy, q)}."""
    rows = []
    for key in D["rows"]:
        m = D["seen"][key]
        if m["mem"] and not with_mem:
            continue
        src = D["inst" if mode == "inst" else "ytd"].get(key, {})
        vals = {}
        for c, k in colmap.items():
            v = (standalone(src, *k) if (mode == "q" and k[1] != 0) else src.get(k if k[1] else (k[0], 4)))
            if v is not None:
                vals[c] = v / SCALE
        if not vals:
            continue
        lab = ko(m["lab"], m["name"], KO, miss)
        rows.append((lab, vals, bool(BOLD.match(local(m["name"]))) and not m["mem"]))
    return rows



def add_button(src, code):
    """리포트 HTML 맨 위(오른쪽)에 'xlsx 다운로드' 버튼 한 번(엑셀이 있을 때만, 스크롤하면 같이 올라감). docs/research·docs/room 의 <코드>/ 폴더 기준 상대 경로."""
    if 'id="es-xlsx"' in src or not os.path.exists(os.path.join(OUT, code + ".xlsx")):
        return src
    btn = ('<style>#es-xlsx-bar{position:relative;z-index:5;text-align:right;padding:12px 16px 0}'
           '#es-xlsx{display:inline-block;padding:8px 14px;border-radius:999px;background:#1d6f42;color:#fff;'
           'font:600 13px/1 system-ui,"Malgun Gothic",sans-serif;text-decoration:none}'
           '#es-xlsx:hover{background:#155a35}@media print{#es-xlsx-bar{display:none}}</style>'
           f'<div id="es-xlsx-bar"><a id="es-xlsx" href="../../data/jp/xlsx/{code}.xlsx" download="{code}.xlsx">xlsx 다운로드</a></div>')
    if re.search(r"<body[^>]*>", src, re.I):
        return re.sub(r"(<body[^>]*>)", lambda m: m.group(1) + btn, src, count=1, flags=re.I)
    return btn + src


def build(code):
    data, used, fsm = collect(code)
    if not used:
        print(code, "결산단신 XBRL 없음")
        return None
    KO, miss = ko_dict(), set()
    name = (json.load(open(PIPE, encoding="utf-8"))["companies"].get(code, {}) or {}).get("n")
    if not name:
        try:
            name = next((w.get("n") for w in json.load(open(os.path.join(ROOT, "docs", "data", "room.json"), encoding="utf-8"))["watch"]
                         if w["c"] == code), None)
        except (OSError, ValueError, KeyError):
            name = None
    name = name or code
    mk = make_model(code, name, data, used, fsm, KO, miss)
    if not mk:
        print(code, "매출 줄을 못 찾음")
        return None
    model, qk, yk, bq, by = mk
    # 양식 그대로 2패스(재무상태표 참조 주소를 먼저 알아야 함)
    wb = openpyxl.Workbook()
    A, qcol, ycol = fmodel.build_annual(wb, model)
    model["_bs_latest_ref"] = fmodel.build_bs(wb, model, A, qcol, ycol)
    wb = openpyxl.Workbook()
    A, qcol, ycol = fmodel.build_annual(wb, model)
    fmodel.build_bs(wb, model, A, qcol, ycol)
    BSS, bcols = model["_bs_sheet"]
    # ── 아랫단: 결산단신 전 항목(FPS 식) ──
    ws = wb["연간"]
    cols = A.cols
    last_col = A.last_col
    cmap = {}
    for c, kind, l in cols:
        if kind == "q":
            cmap[c] = qk[Q_index(l, qk)]
        elif kind == "y":
            cmap[c] = (int(l), 4)
    r = A.r + 1
    blocks = [("손익계산서 · 전 항목", data["PL"], "q", False, "분기 = 3개월 단독"),
              ("포괄손익 · 전 항목", data["CI"], "q", False, "분기 = 3개월 단독"),
              ("부문 · 전 항목", data["SG"], "q", True, "부문 · 항목"),
              ("현금흐름표 · 전 항목", data["CF"], "cum", False, "분기 = 결산기 첫날부터 누계(공시 그대로)")]
    for title, D, mode, mem, note in blocks:
        rows = stmt_rows(D, KO, miss, cmap, mode, mem)
        if rows:
            r = detail_block(ws, r, title, cols, last_col, rows, note)
    ws = wb["재무상태표"]
    bmap = {}
    for c, kind, l in bcols:
        bmap[c] = bq[[qlab(k) for k in bq].index(l)] if kind == "q" else (int(l), 4)
    rows = stmt_rows(data["BS"], KO, miss, bmap, "inst")
    if rows:
        detail_block(ws, BSS.r + 1, "재무상태표 · 전 항목", bcols, BSS.last_col, rows, "분기 말 잔액")
    # 섹터 데이터·출처
    sec = sector_tables(code)
    if sec:
        ws = wb.create_sheet("섹터 데이터")
        ws.sheet_view.showGridLines = False
        ws["A1"] = f"{name} 섹터·사업 데이터"
        ws["A1"].font = Font(name=F, size=20, bold=True)
        ws["A2"] = "2단계 보고서를 만들 때 모은 자료 · 월차 · 출처는 맨 오른쪽"
        ws["A2"].font = Font(name=F, size=9, color="FFFF0000")
        width = max(len(h) for _, h, _, _ in sec) + 2
        ws.column_dimensions["A"].width = 30
        for i in range(2, width):
            ws.column_dimensions[CL(i)].width = 11
        ws.column_dimensions[CL(width)].width = 44
        r = 4
        for t, h, rows_, s in sec:
            r = table_block(ws, r, ko(t, t, KO, miss) if re.search(r"[぀-ヿ]", t) else kt(t), [kt(x) for x in h], rows_, s, width)
    ws = wb.create_sheet("출처")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "출처 · 쓴 결산단신"
    ws["A1"].font = Font(name=F, size=20, bold=True)
    for i, h in enumerate(["공시일", "문서", "PDF"], 1):
        c = ws.cell(3, i, h)
        c.font = Font(name=F, size=10, bold=True)
        c.fill = GRAY
    for i, f in enumerate(used, 4):
        m = re.search(r"(\d{4})年(\d{1,2})月期", f["title"])
        q = re.search(r"第(\d)四半期", f["title"])
        t = (f"{m.group(1)}년 {int(m.group(2))}월 결산기 " if m else "") + (f"{q.group(1)}분기 결산단신" if q else ("중간 결산단신" if "中間" in f["title"] else "연간 결산단신"))
        ws.cell(i, 1, f["date"])
        ws.cell(i, 2, t)
        ws.cell(i, 3, PDF.format(code=code, doc=f["doc"]))
    for col, w in (("A", 12), ("B", 34), ("C", 60)):
        ws.column_dimensions[col].width = w
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(OUT, exist_ok=True)
    out = os.path.join(OUT, code + ".xlsx")
    wb.save(out)
    idx_p = os.path.join(OUT, "index.json")
    try:
        idx = json.load(open(idx_p, encoding="utf-8"))
    except (OSError, ValueError):
        idx = {}
    idx[code] = datetime.now(KST).strftime("%Y-%m-%d %H:%M")
    json.dump(idx, open(idx_p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    if miss:
        try:
            old = json.load(open(MISSING, encoding="utf-8"))
        except (OSError, ValueError):
            old = []
        json.dump(sorted(set(old) | miss), open(MISSING, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{code} 엑셀 → docs/data/jp/xlsx/{code}.xlsx (결산단신 {len(used)}건, 분기 {qlab(qk[0])}~{qlab(qk[-1])}, "
          f"연간 {yk[0]}~{yk[-1]}, 부문 {len(model['segments'])}개, 섹터 표 {len(sec)}개, 번역 필요 {len(miss)}개)")
    return miss


def Q_index(l, qk):
    return [qlab(k) for k in qk].index(l)


if __name__ == "__main__":
    for c in sys.argv[1:]:
        try:
            build(c)
        except Exception as e:
            print(c, "엑셀 실패:", e)
