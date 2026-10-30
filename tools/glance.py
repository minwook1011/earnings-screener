# -*- coding: utf-8 -*-
"""재무제표 한눈에 보기(시제품, 2026-10-29 사용자) — 10년 분기 자료 정리 + 핵심 차트.
  python tools/glance.py <코드> [출력 폴더]
  · 매출·영업이익 10년치는 tools/hist/<코드>.json(위키 jp_financials history=quarterly, 2026-10-30)이 있으면 그것.
    재무상태표·부문·원가는 결산단신 XBRL(tools/xbrl_cache, 없으면 받음) 9년치 — 그보다 옛 분기는 빈칸.
  · 부문(XBRL 세그먼트, 없으면 리서치 <코드>_data.json 의 segments_sales)·수주(orders)·지역(region*)·수출(export*)은 있을 때만.
  · 출력: <코드>_한눈에.xlsx(시트 '재무제표 한눈에 보기' — 위 자료 표, 아래 차트 한 줄 2개) · <코드>_한눈에.html(리포트용 미리보기)
  단위 억 엔. 분기 이름은 달력 분기(3Q26 = 26.7~9).
"""
import glob
import json
import os
import re
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_xlsx as bx
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NQ = 40           # 10년(2026-10-30 — 위키 10년치)
NFIL = 36         # 결산단신 수(8년 + 여유)


def filings_all(code):
    raw = bx.get(bx.LIST.format(code=code))
    out = []
    for it in json.loads(raw or "{}").get("items", []):
        t = it.get("Tdnet", it) if isinstance(it, dict) else {}
        title = t.get("title", "")
        if "決算短信" not in title or "訂正" in title or not t.get("url_xbrl"):
            continue
        m = re.search(r"(\d{18})\.pdf", t.get("document_url", ""))
        if m:
            out.append({"doc": m.group(1), "date": t.get("pubdate", "")[:10], "title": title})
    have = {f["doc"] for f in out}
    out += [f for f in bx.official(code) if f["doc"] not in have]
    return sorted(out, key=lambda f: (f["date"], f["doc"]), reverse=True)[:NFIL]


def cal(k, fsm):
    """(결산기, 분기) → 달력 분기 이름 '3Q26'"""
    fy, q = k
    fy_start = fy - (1 if fsm != 1 else 0)
    m = (fsm - 1 + 3 * q - 1) % 12 + 1
    y = fy_start + (1 if m < fsm else 0)
    return f"{(m - 1) // 3 + 1}Q{y % 100:02d}"


def research_data(code):
    f = glob.glob(os.path.join(bx.RESEARCH, code + " *", code + "_data.json"))
    try:
        return json.load(open(f[0], encoding="utf-8")) if f else {}
    except ValueError:
        return {}


def ext_series(d, keys, labs):
    """리서치 자료의 {labels:[...], 이름:[...]} 표를 우리 분기 칸에 맞춘다. labels 는 '24.7~9' 꼴(끝 달로 맞춤)."""
    for k in keys:
        v = d.get(k)
        if isinstance(v, dict) and isinstance(v.get("labels"), list):
            pos = []
            for s in v["labels"]:
                m = re.match(r"(\d{2})\.(\d{1,2})~(\d{1,2})", str(s))
                pos.append(f"{(int(m.group(3)) - 1) // 3 + 1}Q{m.group(1)}" if m else None)
            out = {}
            for name, arr in v.items():
                if name == "labels" or not isinstance(arr, list) or len(arr) != len(pos):
                    continue
                if not all(isinstance(x, (int, float)) or x is None for x in arr):
                    continue
                row = [None] * len(labs)
                for p, x in zip(pos, arr):
                    if p in labs:
                        row[labs.index(p)] = x
                out[{"total": "합계"}.get(name, name)] = row
            if out:
                return out, v.get("note") or v.get("src") or ""
    return {}, ""


def hist(code, fsm):
    """tools/hist/<코드>.json → {달력 분기: (매출, 영업이익)} 억 엔. 누계 그대로(ytd) 행은 뺀다."""
    f = os.path.join(os.path.dirname(os.path.abspath(__file__)), "hist", code + ".json")
    if not os.path.exists(f):
        return {}, ""
    d = json.load(open(f, encoding="utf-8"))
    out = {}
    for r in d["rows"]:
        m = re.match(r"(\d{4})/\d+期 (\d)Q", r["period"])
        if m and r.get("basis") != "ytd":
            out[cal((int(m.group(1)), int(m.group(2))), fsm)] = (round(r["sales"] / 1e8, 1), round(r["op"] / 1e8, 1))
    return out, d.get("src", "")


def yoy(a):
    return [None if i < 4 or a[i] is None or not a[i - 4] or a[i - 4] <= 0 else round((a[i] / a[i - 4] - 1) * 100, 1) for i in range(len(a))]


# 총 차입금 = 이자 붙는 부채 전부(차입금·CP·사채·전환사채·리스부채). 한 줄 = 한 항목, 줄 안은 같은 뜻의 다른 이름(앞 것 우선)
IBD = [["ShortTermLoansPayable", "ShortTermBorrowings"], ["CommercialPapersLiabilities", "CommercialPapersCLIFRS"],
       ["ShortTermBondsPayable"], ["CurrentPortionOfLongTermLoansPayable", "CurrentPortionOfLongTermBorrowingsCL"],
       ["CurrentPortionOfBonds"], ["CurrentPortionOfConvertibleBonds"],
       ["LongTermLoansPayable", "LongTermBorrowings"], ["BondsPayable"],
       ["ConvertibleBondTypeBondsWithSubscriptionRightsToShares", "ConvertibleBondTypeBondsWithSubscriptionRightsToSharesNCL"],
       ["LeaseObligationsCL", "LeaseLiabilitiesCL", "LeaseLiabilitiesCLIFRS"],
       ["LeaseObligationsNCL", "LeaseLiabilitiesNCL", "LeaseLiabilitiesNCLIFRS"],
       ["BondsAndBorrowingsCLIFRS", "BorrowingsCLIFRS"], ["BondsAndBorrowingsNCLIFRS", "BorrowingsNCLIFRS"]]


def gather(code):
    orig, bx.filings = bx.filings, filings_all
    try:
        data, used, fsm = bx.collect(code)
    finally:
        bx.filings = orig
    PL, BS = data["PL"], data["BS"]
    keys = bx.qseq(max(k for k in bx.pick(PL, bx.PLMAP["revenue"])), NQ)
    labs = [cal(k, fsm) for k in keys]

    def flow(names, D=PL):
        y = bx.pick(D, names)
        return [None if bx.standalone(y, *k) is None else round(bx.standalone(y, *k) / 1e8, 1) for k in keys]

    def stock(names):
        y = bx.pick(BS, names, "inst")
        return [None if y.get(k) is None else round(y[k] / 1e8, 1) for k in keys]

    rev, op, cogs = flow(bx.PLMAP["revenue"]), flow(bx.PLMAP["op"]), flow(bx.PLMAP["cogs"])
    H, hsrc = hist(code, fsm)      # 10년치가 있으면 매출·영업이익은 그것으로(결산단신 XBRL 은 9년까지만 받음)
    for i, l in enumerate(labs):
        if l in H:
            rev[i], op[i] = H[l]
    ca, cl, tl = stock(bx.BSMAP["current_assets"]), stock(bx.BSMAP["current_liab"]), stock(bx.BSMAP["total_liab"])
    eq, inv, rec = stock(bx.BSMAP["total_equity"]), stock(bx.BSMAP["inventory"]), stock(bx.BSMAP["receivables"])
    cash, pay = stock(bx.BSMAP["cash"]), stock(bx.BSMAP["payables"])
    if all(x is None for x in inv):      # 재고를 제품·재공품·원재료로 나눠 내는 회사
        parts = [stock([n]) for n in bx.INV_PARTS]
        inv = [round(sum(x for x in t if x is not None), 1) if any(x is not None for x in t) else None for t in zip(*parts)]
    debt = [sum(x for x in t if x is not None) if any(x is not None for x in t) else (0 if e is not None else None)
            for t, e in zip(zip(*[stock(g) for g in IBD]), eq)]   # 재무상태표가 없는 분기는 빈칸(차입금 없는 회사만 0)
    div = lambda a, b, m=100: [None if x is None or not y else round(x / y * m, 1) for x, y in zip(a, b)]
    T = {"매출액": rev, "매출액 YoY": yoy(rev), "영업이익": op, "영업이익률": div(op, rev), "영업이익 YoY": yoy(op),
         "유동자산": ca, "유동부채": cl, "부채총계": tl, "자기자본": eq,
         "유동비율": div(ca, cl), "부채비율": div(tl, eq),
         "순차입금/자기자본": [None if e is None or c is None else round((d - c) / e * 100, 1) for d, c, e in zip(debt, cash, eq)],
         "재고자산": inv, "재고자산 회전일수": [None if i is None or not c else round(i / c * 91, 0) for i, c in zip(inv, cogs)],
         "매출채권 회전일수": [None if r is None or not s else round(r / s * 91, 0) for r, s in zip(rec, rev)],
         "매입채무 회전일수": [None if r is None or not s else round(r / s * 91, 0) for r, s in zip(pay, cogs)]}
    T["현금전환주기(CCC)"] = [None if None in t else t[0] + t[1] - t[2] for t in zip(T["재고자산 회전일수"], T["매출채권 회전일수"], T["매입채무 회전일수"])]
    # 부문: XBRL 세그먼트 매출(외부 고객) → 없으면 리서치 자료
    seg, segsrc = {}, ""
    SG = data["SG"]
    for k in SG["rows"]:
        m = SG["seen"][k]
        if m["mem"] and re.search(r"NetSales|Revenue", m["name"]) and re.search(r"外部顧客|External", m["lab"] + m["name"]):
            mem = m["mem"][-1].split(":")[-1]
            if mem in bx.STDMEM:
                continue
            y = SG["ytd"].get(k, {})
            row = [None if bx.standalone(y, *kk) is None else round(bx.standalone(y, *kk) / 1e8, 1) for kk in keys]
            if sum(x is not None for x in row) >= 4:
                seg[m["lab"].split(" · ")[0]] = row
    if seg:
        segsrc = "결산단신 부문 정보(외부 고객 매출)"
    rd = research_data(code)
    if not seg:
        seg, segsrc = ext_series(rd, ["segments_sales", "segment_sales", "segments"], labs)
    kpi, ksrc = kpi_file(code, labs)          # 리서치 폴더 <코드>_kpi.json — 회사 설명회 자료에서 뽑은 8년치
    if kpi.get("부문별 매출"):
        seg, segsrc = kpi["부문별 매출"], ksrc
    region, rsrc = (kpi["지역별 매출"], ksrc) if kpi.get("지역별 매출") else ext_series(rd, ["region_sales", "sales_region", "regions"], labs)
    orders, osrc = (kpi["부문별 수주"], ksrc) if kpi.get("부문별 수주") else ext_series(rd, ["orders", "order"], labs)
    exports, esrc = trade_series(code, labs)
    T["총 차입금"], T["현금"] = debt, cash
    SRC = {}
    for grp, (d, src) in (("부문별 매출", (seg, segsrc)), ("지역별 매출", (region, rsrc)), ("수주", (orders, osrc)), ("수출", (exports, esrc))):
        for name, arr in d.items():
            T[f"{grp} · {name}"] = arr
            SRC[f"{grp} · {name}"] = src
        if grp == "수주" and d:
            parts = {k: v for k, v in d.items() if k != "합계"}
            tot = d.get("합계") or [sum(x for x in t if x is not None) if any(x is not None for x in t) else None for t in zip(*parts.values())]
            T["수주 · 합계"], T["수주 YoY"] = tot, yoy(tot)
            SRC["수주 YoY"] = src
        if grp == "수출":
            for name, arr in d.items():
                T[f"수출 YoY · {name}"] = yoy(arr)
                SRC[f"수출 YoY · {name}"] = src
    return {"code": code, "labs": labs, "T": T, "SRC": SRC, "n_fil": len(used), "hsrc": hsrc,
            "first": used[-1]["date"] if used else "", "last": used[0]["date"] if used else ""}


def kpi_file(code, labs):
    f = glob.glob(os.path.join(bx.RESEARCH, code + " *", code + "_kpi.json"))
    if not f:
        return {}, ""
    d = json.load(open(f[0], encoding="utf-8"))
    out = {}
    for grp, v in d.items():
        if grp.startswith("_") or not isinstance(v, dict):
            continue
        rows = {n: [s.get(q) for q in labs] for n, s in v.items() if isinstance(s, dict)}
        out[grp] = {n: r for n, r in rows.items() if sum(x is not None for x in r) >= 4}
    return out, d.get("_src", "")


def trade_series(code, labs):
    """리서치 폴더 lead_input.json 의 월별 수출(천엔, 일본 무역통계) → 분기 합(억 엔)"""
    f = glob.glob(os.path.join(bx.RESEARCH, code + " *", "lead_input.json"))
    if not f:
        return {}, ""
    ind = json.load(open(f[0], encoding="utf-8")).get("indicators", {})
    out = {}
    for name, v in ind.items():
        arr = v.get("data") if isinstance(v, dict) else v
        if not isinstance(arr, list) or not re.search(r"수출|수입", name):
            continue
        nm = re.sub(r"\s*yoy\w*$", "", name)
        if nm in out:
            continue
        q = {}
        for ym, val in arr:
            y, m = int(ym[:4]), int(ym[5:7])
            q.setdefault(f"{(m - 1) // 3 + 1}Q{y % 100:02d}", []).append(val)
        out[nm] = [round(sum(q[l]) / 1e5, 1) if l in q and len(q[l]) == 3 else None for l in labs]
    return out, "일본 무역통계(재무성, 월별 → 분기 합)"


# ---------- 차트 목록(엑셀·HTML 같이 씀) ----------
def charts(G):
    """사용자 지정(2026-10-29): 매출 · 매출 YoY+영업이익 YoY · 부문별 · 지역별 · 수주+수주 YoY · 영업이익+OPM · 수출입 · 재무 안정성.
    시리즈 이름 = 표의 행 이름(엑셀이 그 행을 가리킴)."""
    T, SRC, C = G["T"], G["SRC"], []
    pick = lambda pre: {k: v for k, v in T.items() if k.startswith(pre + " · ") and not k.endswith("· 합계")}
    src = lambda d: next((SRC[k] for k in d if SRC.get(k)), "")
    # YoY 는 따로 두지 않고 매출·영업이익 차트에 꺾은선으로 겹친다(2026-10-30 사용자)
    C.append({"t": "매출액 (억 엔) · YoY (%)", "bar": {"매출액": T["매출액"]}, "line2": {"매출액 YoY": T["매출액 YoY"]}})
    for grp in ("부문별 매출", "지역별 매출"):
        d = pick(grp)
        if d:
            C.append({"t": grp + " (억 엔)", "line": d, "src": src(d)})
    d = pick("수주")
    if d:
        C.append({"t": "수주 (억 엔) · 수주 YoY (%)", "stack": d, "line2": {"수주 YoY": T["수주 YoY"]}, "src": src(d)})
    C.append({"t": "영업이익 (억 엔) · OPM·YoY (%)", "bar": {"영업이익": T["영업이익"]},
              "line2": {"영업이익률": T["영업이익률"]}, "line3": {"영업이익 YoY": T["영업이익 YoY"]}})   # YoY 는 범위가 커서 축 따로
    d = pick("수출")
    if d:
        C.append({"t": "관련 수출 (억 엔) · YoY (%)", "bar": d, "line2": pick("수출 YoY"), "src": src(d)})
    nodebt = not any(T["총 차입금"])   # 차입금·사채·리스부채가 하나도 없는 회사(화낙 등) — 막대가 안 보이는 이유를 적어 둠
    C.append({"t": "재무 안정성 — 총 차입금·현금 (억 엔) · 부채비율 (%)", "bar": {"총 차입금": T["총 차입금"], "현금": T["현금"]},
              "line2": {"부채비율": T["부채비율"]}, "src": "총 차입금 0 — 차입금·사채·리스부채 없음" if nodebt else ""})
    return C


# ---------- 엑셀 ----------
BLUE, ORANGE, GRAY, GOLD, GREEN, NAVY = "4472C4", "ED7D31", "A5A5A5", "FFC000", "70AD47", "264478"
PAL = [BLUE, ORANGE, GRAY, GOLD, GREEN, NAVY, "9E480E", "636363"]


def xlsx(G, path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, G)
    wb.save(path)


def sheet(wb, G):
    """wb 에 '재무제표 한눈에 보기' 시트를 더한다(표준 재무제표 엑셀에도 이걸로 붙임)."""
    ws = wb.create_sheet("재무제표 한눈에 보기")
    labs, n = G["labs"], len(G["labs"])
    F, B = Font(name="맑은 고딕", size=9), Font(name="맑은 고딕", size=9, bold=True)
    ws["A1"] = f"{G['code']} 재무제표 한눈에 보기 — 최근 {n}개 분기(달력 분기, 3Q26 = 26.7~9) · 단위 억 엔"
    ws["A1"].font = Font(name="맑은 고딕", size=12, bold=True)
    ws["A2"] = f"출처: 매출·영업이익 {G['hsrc'] or '결산단신 XBRL'} · 재무상태표 등 결산단신 XBRL {G['n_fil']}건({G['first']}~{G['last']}). 부문·지역·수주·수출은 있을 때만(출처는 줄 끝). 총 차입금 = 이자 붙는 부채 전부(단기·장기 차입금, CP, 사채·전환사채, 리스부채)"
    ws["A2"].font = Font(name="맑은 고딕", size=8, color="666666")
    hdr = 4
    ws.cell(hdr, 1, "항목").font = B
    for j, l in enumerate(labs):
        c = ws.cell(hdr, 2 + j, l)
        c.font, c.alignment, c.fill = B, Alignment(horizontal="center"), PatternFill("solid", fgColor="D9E1F2")
    ws.cell(hdr, 2 + n, "출처").font = B
    rows, r = {}, hdr + 1
    for nm, arr in G["T"].items():
        ws.cell(r, 1, nm).font = F
        pct = re.search(r"YoY|비율|이익률|자기자본$", nm)
        for j, v in enumerate(arr):
            c = ws.cell(r, 2 + j, v)
            c.font = F
            c.number_format = '0.0"%"' if pct else '#,##0'
        if G["SRC"].get(nm):
            ws.cell(r, 2 + n, G["SRC"][nm]).font = Font(name="맑은 고딕", size=8, color="666666")
        rows[nm] = r
        r += 1
    ws.column_dimensions["A"].width = 22
    for j in range(n):
        ws.column_dimensions[get_column_letter(2 + j)].width = 7.5
    ws.freeze_panes = ws.cell(hdr + 1, 2)

    def ref(nm):
        rr = rows[nm]
        return Reference(ws, min_col=1, max_col=1 + n, min_row=rr, max_row=rr)
    cats = Reference(ws, min_col=2, max_col=1 + n, min_row=hdr, max_row=hdr)
    W, H, top = 24, 9, r + 2      # 차트 한 개 24cm × 9cm, 한 줄에 2개
    for i, ch in enumerate(charts(G)):
        main = None
        if ch.get("bar") or ch.get("stack"):
            main = BarChart()
            main.type = "col"
            if ch.get("stack"):
                main.grouping, main.overlap = "stacked", 100
            for k in (ch.get("bar") or ch.get("stack")):
                main.add_data(ref(k), titles_from_data=True, from_rows=True)
            main.gapWidth = 40
        if ch.get("line") and main is None:
            main = LineChart()
            for k in ch["line"]:
                main.add_data(ref(k), titles_from_data=True, from_rows=True)
        if ch.get("line2"):
            l2 = LineChart()
            for k in list(ch["line2"]) + list(ch.get("line3") or {}):
                l2.add_data(ref(k), titles_from_data=True, from_rows=True)
            l2.y_axis.axId, l2.y_axis.crosses = 200, "max"
            l2.y_axis.title = "%" if "회전" not in ch["t"] else "일"
            l2.y_axis.majorGridlines = None
            l2.y_axis.delete = False
            l2.dataLabels = DataLabelList()
            l2.dataLabels.showVal, l2.dataLabels.position = True, "t"
            l2.dataLabels.showSerName = l2.dataLabels.showCatName = l2.dataLabels.showLegendKey = False
            l2.dataLabels.numFmt = '0"%"' if "회전" not in ch["t"] else '0'
            for s in l2.series:
                s.marker.symbol, s.marker.size, s.smooth = "circle", 5, False
        if main is None:
            continue
        main.set_categories(cats)
        main.title, main.width, main.height = ch["t"], W, H
        if main.legend is not None:
            main.legend.position = "t"
        main.y_axis.majorGridlines = None
        main.x_axis.delete = main.y_axis.delete = False
        main.x_axis.tickLblPos = "low"
        for j, s in enumerate(main.series):
            col = PAL[j % len(PAL)]
            if isinstance(main, BarChart):
                s.graphicalProperties.solidFill = col
                s.graphicalProperties.line.solidFill = col
            else:
                s.graphicalProperties.line.solidFill, s.graphicalProperties.line.width = col, 22000
                s.marker.symbol, s.marker.size = "circle", 5
                s.marker.graphicalProperties.solidFill = s.marker.graphicalProperties.line.solidFill = col
                s.smooth = False
        if (len(main.series) == 1 or ch.get("line2")) and not ch.get("stack"):   # 쌓은 막대 + outEnd 이름표는 엑셀이 못 엶
            main.dataLabels = DataLabelList()
            main.dataLabels.showVal = True
            main.dataLabels.showSerName = main.dataLabels.showCatName = main.dataLabels.showLegendKey = False
            main.dataLabels.position = "outEnd" if isinstance(main, BarChart) else "t"
            main.dataLabels.numFmt = '0"%"' if ch.get("pct") else '#,##0'
        main.varyColors = False
        if len(main.series) == 1 and not ch.get("line2"):
            main.legend = None
        if ch.get("line2"):
            for j, s in enumerate(l2.series):
                col = PAL[(j + len(main.series)) % len(PAL)]
                s.graphicalProperties.line.solidFill = col
                s.marker.graphicalProperties.solidFill = s.marker.graphicalProperties.line.solidFill = col
            main += l2
        cell = f"{'A' if i % 2 == 0 else get_column_letter(2 + n // 2 + 1)}{top + (i // 2) * 19}"
        ws.add_chart(main, cell)


# ---------- 리포트용 HTML(한 줄 2개) ----------
HTML = """<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex"><title>{code} 재무제표 한눈에 보기</title>
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root{{--bg:#f4f7fb;--card:#fff;--ink:#1f2937;--mute:#6b7280;--line:#e5e7eb}}
@media (prefers-color-scheme:dark){{:root{{--bg:#0f172a;--card:#111827;--ink:#e5e7eb;--mute:#9ca3af;--line:#1f2937}}}}
body{{margin:0;background:var(--bg);color:var(--ink);font-family:"Pretendard","Malgun Gothic",sans-serif}}
.wrap{{max-width:1280px;margin:0 auto;padding:20px 16px}} h2{{font-size:20px;margin:0 0 4px}} .sub{{color:var(--mute);font-size:12px;margin-bottom:14px}}
.grid{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px}} @media (max-width:760px){{.grid{{grid-template-columns:1fr}}}}
.c{{min-width:0;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 12px 6px}} .c h3{{font-size:14px;margin:0 0 6px}}
.c .s{{font-size:11px;color:var(--mute);margin-top:4px}} .cv{{position:relative;height:300px}}
</style></head><body><div class="wrap"><h2>재무제표 한눈에 보기 — {code}</h2>
<div class="sub">최근 {n}개 분기(달력 분기, 3Q26 = 26.7~9) · 억 엔 · 매출·영업이익 10년치(위키 jp_financials) · 재무상태표 등 결산단신 XBRL {nf}건({first}~{last}), 그 전 분기는 빈칸. 부문·수주·지역·수출은 있을 때만.</div>
<div class="grid" id="g"></div></div>
<script>
const L={labs},C={charts};
const PAL=["#4472C4","#ED7D31","#A5A5A5","#FFC000","#70AD47","#264478","#9E480E","#636363"];
const dark=matchMedia("(prefers-color-scheme: dark)").matches, ink=dark?"#e5e7eb":"#1f2937", grid=dark?"#1f2937":"#eef0f3";
Chart.defaults.font.family='"Pretendard","Malgun Gothic",sans-serif';Chart.defaults.color=ink;Chart.defaults.font.size=10;
const lab={{id:"lab",afterDatasetsDraw(c){{const x=c.ctx;x.save();x.font="bold 9px sans-serif";x.textAlign="center";
 c.data.datasets.forEach((d,i)=>{{if(!d.lab)return;const m=c.getDatasetMeta(i);if(m.hidden)return;x.fillStyle=d.borderColor;
 m.data.forEach((p,j)=>{{const v=d.data[j];if(v==null||j%d.lab)return;x.fillText((d.pct?Math.round(v)+"%":Math.round(v).toLocaleString()),p.x,p.y-6)}})}});x.restore()}}}};
C.forEach((ch,i)=>{{const el=document.createElement("div");el.className="c";el.innerHTML="<h3>"+ch.t+"</h3><div class='cv'><canvas></canvas></div>"+(ch.src?"<div class='s'>출처: "+ch.src+"</div>":"");
 document.getElementById("g").appendChild(el);const ds=[];let k=0;
 const step=Math.max(1,Math.ceil(L.length/16));
 const add=(o,type,axis,extra)=>Object.entries(o||{{}}).forEach(([n,a])=>{{const c=PAL[k++%PAL.length];
  ds.push(Object.assign({{label:(n.indexOf("수출 YoY")==0?n.split(" · ").pop()+" YoY":n.split(" · ").pop()),data:a,type,yAxisID:axis,backgroundColor:c,borderColor:c,borderWidth:type=="line"?2:0,pointRadius:type=="line"?3:0,spanGaps:true,lab:(Object.keys(o).length==1||axis=="y2")?step:0,pct:axis!="y"&&!/회전/.test(ch.t)||ch.pct}},extra||{{}}))}});
 add(ch.bar,"bar","y");add(ch.stack,"bar","y",{{stack:"s"}});add(ch.line,"line","y");add(ch.line2,"line","y2",{{order:-1}});add(ch.line3,"line","y3",{{order:-2,borderDash:[4,3]}});
 const sc={{x:{{grid:{{display:false}},stacked:!!ch.stack,ticks:{{maxRotation:90,minRotation:90,autoSkip:false}}}},y:{{grid:{{color:grid}},stacked:!!ch.stack,ticks:{{callback:v=>ch.pct?v+"%":v.toLocaleString()}}}}}};
 if(ch.line2)sc.y2={{position:"right",grid:{{display:false}},ticks:{{callback:v=>/회전/.test(ch.t)?v+"일":v+"%"}}}};if(ch.line3)sc.y3={{position:"right",grid:{{display:false}},ticks:{{callback:v=>v+"%"}}}};
 new Chart(el.querySelector("canvas"),{{type:"bar",data:{{labels:L,datasets:ds}},plugins:[lab],options:{{maintainAspectRatio:false,interaction:{{mode:"index",intersect:false}},plugins:{{legend:{{position:"bottom",labels:{{boxWidth:10}}}}}},scales:sc}}}})}});
</script></body></html>"""


def html(G, path):
    open(path, "w", encoding="utf-8").write(HTML.format(code=G["code"], n=len(G["labs"]), nf=G["n_fil"], first=G["first"], last=G["last"],
                                                     labs=json.dumps(G["labs"]), charts=json.dumps(charts(G), ensure_ascii=False)))


# ---------- 2단계 리포트 안 토글(2026-10-30 사용자 — 스킬에 추가) ----------
FRAG = """<details class="tg gl" id="es-glance"><summary><b>재무제표 한눈에 보기 — 최근 {n}개 분기 핵심 차트</b></summary><div class="tgb">
<style>#es-glance .glx-sub{{color:#6b7280;font-size:12px;margin:0 0 10px}}#es-glance .glx-g{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}}
@media (max-width:760px){{#es-glance .glx-g{{grid-template-columns:1fr}}}}#es-glance .glx-c{{min-width:0;background:#fff;border:1px solid #e5e7eb;border-radius:10px;padding:10px 10px 4px}}
#es-glance .glx-c h4{{font-size:13px;margin:0 0 6px}}#es-glance .glx-s{{font-size:11px;color:#6b7280;margin-top:4px}}#es-glance .glx-v{{position:relative;height:280px}}</style>
<p class="glx-sub">달력 분기(3Q26 = 26.7~9) · 억 엔 · {src} · 부문·지역·수주·수출은 있을 때만 · 엑셀 '재무제표 한눈에 보기' 시트에 같은 표·차트</p>
<div class="glx-g"></div></div></details>
<script>(function(){{const L={labs},C={charts};const box=document.getElementById("es-glance");let done=false;
function draw(){{if(done)return;done=true;const PAL=["#4472C4","#ED7D31","#A5A5A5","#FFC000","#70AD47","#264478","#9E480E","#636363"];
Chart.defaults.font.family='"Pretendard","Malgun Gothic",sans-serif';
const lab={{id:"glab",afterDatasetsDraw(c){{const x=c.ctx;x.save();x.font="bold 9px sans-serif";x.textAlign="center";
 c.data.datasets.forEach((d,i)=>{{if(!d.lab)return;const m=c.getDatasetMeta(i);if(m.hidden)return;x.fillStyle=d.borderColor;
 m.data.forEach((p,j)=>{{const v=d.data[j];if(v==null||j%d.lab)return;x.fillText((d.pct?Math.round(v)+"%":Math.round(v).toLocaleString()),p.x,p.y-6)}})}});x.restore()}}}};
C.forEach(ch=>{{const el=document.createElement("div");el.className="glx-c";el.innerHTML="<h4>"+ch.t+"</h4><div class='glx-v'><canvas></canvas></div>"+(ch.src?"<div class='glx-s'>"+(/^총 차입금/.test(ch.src)?"":"출처: ")+ch.src+"</div>":"");
 box.querySelector(".glx-g").appendChild(el);const ds=[];let k=0;const step=Math.max(1,Math.ceil(L.length/12));
 const add=(o,type,axis,extra)=>Object.entries(o||{{}}).forEach(([n,a])=>{{const c=PAL[k++%PAL.length];
  ds.push(Object.assign({{label:(n.indexOf("수출 YoY")==0?n.split(" · ").pop()+" YoY":n.split(" · ").pop()),data:a,type,yAxisID:axis,backgroundColor:c,borderColor:c,borderWidth:type=="line"?2:0,pointRadius:type=="line"?2.5:0,spanGaps:true,lab:(Object.keys(o).length==1&&axis=="y")?step:0,pct:axis!="y"||ch.pct}},extra||{{}}))}});
 add(ch.bar,"bar","y");add(ch.stack,"bar","y",{{stack:"s"}});add(ch.line,"line","y");add(ch.line2,"line","y2",{{order:-1}});add(ch.line3,"line","y3",{{order:-2,borderDash:[4,3]}});
 const sc={{x:{{grid:{{display:false}},stacked:!!ch.stack,ticks:{{maxRotation:90,minRotation:90,autoSkip:false,font:{{size:9}}}}}},y:{{grid:{{color:"#eef0f3"}},stacked:!!ch.stack,ticks:{{callback:v=>ch.pct?v+"%":v.toLocaleString()}}}}}};
 if(ch.line2)sc.y2={{position:"right",grid:{{display:false}},ticks:{{callback:v=>v+"%"}}}};if(ch.line3)sc.y3={{position:"right",grid:{{display:false}},ticks:{{callback:v=>v+"%"}}}};
 new Chart(el.querySelector("canvas"),{{type:"bar",data:{{labels:L,datasets:ds}},plugins:[lab],options:{{maintainAspectRatio:false,interaction:{{mode:"index",intersect:false}},plugins:{{legend:{{position:"bottom",labels:{{boxWidth:10,font:{{size:10}}}}}}}},scales:sc}}}})}})}}
function go(){{if(window.Chart)return draw();const sc=document.createElement("script");sc.src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js";sc.onload=draw;document.head.appendChild(sc)}}
box.addEventListener("toggle",()=>{{if(box.open)go()}});if(box.open)go()}})();</script>"""


def fragment(G):
    src = ("매출·영업이익 " + G["hsrc"].split("(")[0] + " · " if G.get("hsrc") else "") + f"결산단신 XBRL {G['n_fil']}건({G['first']}~{G['last']})"
    return FRAG.format(n=len(G["labs"]), src=src, labs=json.dumps(G["labs"]), charts=json.dumps(charts(G), ensure_ascii=False))


def insert(code, src):
    """리포트 HTML 에 토글을 넣는다 — 재무제표 표 바로 아래(첫 토글 앞). 이미 있으면 바꾼다. 실패하면 그대로."""
    G = gather(code)
    fr = fragment(G)
    src = re.sub(r'<details class="tg gl" id="es-glance">.*?</details>\s*<script>\(function\(\)\{const L=.*?</script>', "", src, flags=re.S)
    m = re.search(r'<details class="tg (?!prof)', src)
    if not m:
        return src
    return src[:m.start()] + fr + src[m.start():]


if __name__ == "__main__":
    code = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else "."
    G = gather(code)
    os.makedirs(out, exist_ok=True)
    xlsx(G, os.path.join(out, f"{code}_한눈에.xlsx"))
    html(G, os.path.join(out, f"{code}_한눈에.html"))
    print(code, len(G["labs"]), "분기", G["labs"][0], "~", G["labs"][-1])
