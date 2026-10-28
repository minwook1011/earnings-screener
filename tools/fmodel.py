#!/usr/bin/env python3
"""재무모델링 엑셀 생성기 — model.json → {기업}_재무모델.xlsx
(사용자 「재무모델링」 스킬의 scripts/build_model.py 사본 — 테크윙 양식 고정. 바꾼 곳 2개:
 meta.labels 로 기타이익·기타손실·금융수익·금융원가 이름을 바꿀 수 있음(일본 기준 = 특별이익·특별손실·영업외수익·영업외비용),
 build_bs 가 d["_bs_sheet"] 에 시트·열 정보를 남김(아래에 상세 표를 잇기 위해))

시트 2개를 만든다.
  1) 연간      : 12분기 + 10개년 실적, 5개년 추정(MW) 칸, 컨센서스(E) 칸, 순이익 추정·목표주가·자산가치
  2) 재무상태표 : 최근 5분기 + 5개년 BS, 현금흐름 요약, 주요 재무 비율(버핏·펀더멘털 체크)

사용법:
  python build_model.py model.json -o 테크윙_재무모델.xlsx

JSON 규칙 (자세한 스키마는 references/json_schema.md):
  - 금액은 원 통화 단위 그대로(원·USD·円) 넣는다. meta.scale 로 나눠서(억원=1e8, $M=1e6) 시트에 표시한다.
  - EPS·주가·주식수는 나누지 않는다.
  - 값이 없으면 null. 빈칸으로 남고 수식도 빈칸을 돌려준다 (0으로 채우지 않는다).
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

# ── 디자인 토큰 (예시 「연간」 시트의 회색 톤을 그대로 따른다) ─────────────────
FONT = "맑은 고딕"
GRAY_HDR = "D9D9D9"      # 열 헤더 행 배경 (데이터 행은 흰색)
GRAY_SEC = "808080"      # 섹션 제목 배경 (흰 글씨)
GRAY_SUB = "A6A6A6"      # 소섹션 제목 배경
GRAY_TXT = "808080"      # 비중·% of Sales 같은 보조 행 글씨
INPUT_BLUE = "0000FF"    # 추정(MW)·가정 입력칸
RED = "FF0000"

NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;-0.0%;"-"'
MULT = '0.0"x";-0.0"x";"-"'
DAYS = '0"일";-0"일";"-"'
YRS = '0.0"년";-0.0"년";"-"'
WON = '#,##0'

thin = Side(style="thin", color="7F7F7F")

DEFAULT_TAX = {"KR": 0.22, "US": 0.21, "JP": 0.30}


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def font(bold=False, color="000000", size=11, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)


# ── 공통 헬퍼 ─────────────────────────────────────────────────────────────
def blank_safe(expr, *cells):
    """참조 셀 중 하나라도 비어 있으면 빈칸을 돌려주는 수식."""
    if not cells:
        return f"=IFERROR({expr},\"\")"
    cond = ",".join(f'{c}=""' for c in cells)
    return f'=IF(OR({cond}),"",IFERROR({expr},""))'


class Sheet:
    """행 단위로 쓰면서 key → 행 번호를 기억한다."""

    def __init__(self, ws, cols, label_width=26):
        self.ws = ws
        self.cols = cols          # [(col_idx, kind, label)] kind: q / y / mw / cs
        self.rows = {}
        self.r = 4
        ws.column_dimensions["A"].width = label_width
        ws.column_dimensions["B"].width = 14
        for c, kind, _ in cols:
            ws.column_dimensions[CL(c)].width = 11 if kind != "cs" else 12
        self.last_col = cols[-1][0] if cols else 2

    def ref(self, key, col):
        return f"{CL(col)}{self.rows[key]}"

    def style_row(self, r, style):
        ws = self.ws
        for c in range(1, self.last_col + 1):
            cell = ws.cell(r, c)
            cell.alignment = CENTER
            if style == "section":
                cell.fill = fill(GRAY_SEC)
                cell.font = font(True, "FFFFFF")
            elif style == "sub":
                cell.fill = fill(GRAY_SUB)
                cell.font = font(True, "FFFFFF")
            elif style == "key":
                cell.font = font(True)
            elif style == "detail":
                cell.font = font(False)
            elif style == "ratio":
                cell.font = font(False, GRAY_TXT)
            else:
                cell.font = font(False)
        ws.cell(r, 1).alignment = LEFT if style in ("ratio",) else CENTER

    def section(self, title, note=None, sub=False):
        r = self.r
        self.ws.cell(r, 1, title)
        if note:
            self.ws.cell(r, 2, note)
        self.style_row(r, "sub" if sub else "section")
        self.r += 1
        return r

    def row(self, key, label, cell_fn, style="plain", fmt=NUM, note=None,
            input_kinds=()):
        """cell_fn(col, kind, label) → 값/수식/None. input_kinds 의 칸은 파란 글씨."""
        ws, r = self.ws, self.r
        ws.cell(r, 1, label)
        if note is not None:
            ws.cell(r, 2, note)
        self.style_row(r, style)
        self.rows[key] = r
        for c, kind, plabel in self.cols:
            v = cell_fn(c, kind, plabel)
            cell = ws.cell(r, c)
            if v is not None:
                cell.value = v
            cell.number_format = fmt
            if kind in input_kinds:
                cell.font = font(style == "key", INPUT_BLUE)
        ws.cell(r, 2).font = font(False, GRAY_TXT, 9)
        self.r += 1
        return r

    def gap(self):
        self.r += 1


def scaled(v, scale):
    if v is None:
        return None
    return v / scale


def series(block, key, n):
    vals = (block or {}).get(key)
    if vals is None:
        return [None] * n
    vals = list(vals) + [None] * (n - len(vals))
    return vals[:n]


def group_cols(ws, cols):
    """분기 열을 엑셀 그룹(1·2 단추)으로 묶어 접었다 펼 수 있게 한다. 기본은 펼친 상태."""
    if not cols:
        return
    for c in cols:
        ws.column_dimensions[CL(c)].outline_level = 1
    ws.sheet_properties.outlinePr.summaryRight = True


def print_setup(ws):
    """가로 A3, 가로 한 페이지에 맞춤."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"


# ── 시트 1: 연간 ─────────────────────────────────────────────────────────
def build_annual(wb, d):
    meta, scale = d["meta"], d["meta"]["scale"]
    mkt = meta.get("market", "KR")
    Q, Y = d["quarterly"], d["annual"]
    qlab, ylab = Q["labels"], Y["labels"]
    fc = d.get("forecast", {})
    mw_years = fc.get("mw_years", [])
    cs = fc.get("consensus") or {}
    cslab = cs.get("labels", [])

    ws = wb.active
    ws.title = "연간"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85

    # 열 배치: A 항목 / B 참고 / 분기 12 / 연간 10 / MW 5 / 컨센서스
    cols, c = [], 3
    for l in qlab:
        cols.append((c, "q", l)); c += 1
    for l in ylab:
        cols.append((c, "y", str(l))); c += 1
    for l in mw_years:
        cols.append((c, "mw", f"{l}(MW)")); c += 1
    for l in cslab:
        cols.append((c, "cs", l)); c += 1
    S = Sheet(ws, cols)
    qcol = {l: c for c, k, l in cols if k == "q"}
    ycol = {l: c for c, k, l in cols if k == "y"}
    qcols = [c for c, k, _ in cols if k == "q"]
    ycols = [c for c, k, _ in cols if k == "y"]
    mwcols = [c for c, k, _ in cols if k == "mw"]
    cscols = [c for c, k, _ in cols if k == "cs"]
    last_y = ycols[-1] if ycols else None

    # 제목·헤더
    ws.cell(1, 1, f"{meta['company']} Model").font = font(True, size=20)
    ws.row_dimensions[1].height = 30
    ws.cell(2, 1, "모를 땐 보수적으로").font = font(False, RED)
    ws.cell(2, 2, f"(단위: {meta['unit_label']})").font = font(False, GRAY_TXT, 9)
    if meta.get("fiscal_note"):
        ws.cell(2, 3, meta["fiscal_note"]).font = font(False, GRAY_TXT, 9)
    if cscols:
        ws.cell(2, cscols[0], "(컨센서스)").font = font(True)
        ws.cell(2, cscols[0]).alignment = CENTER
    if mwcols:
        ws.cell(2, mwcols[0], "(추정: 파란 글씨 = 직접 입력)").font = font(False, INPUT_BLUE, 9)
    ws.cell(3, 1, "항목"); ws.cell(3, 2, "참고")
    for c, k, l in cols:
        ws.cell(3, c, l)
    for c in range(1, S.last_col + 1):
        cell = ws.cell(3, c)
        cell.fill = fill(GRAY_HDR)
        cell.font = font(True, size=10)
        cell.alignment = CENTER
        cell.border = Border(top=thin, bottom=thin)

    def prev_col(col, kind):
        """YoY 비교 대상 열."""
        if kind == "q":
            i = qcols.index(col)
            return qcols[i - 4] if i >= 4 else None
        if kind == "y":
            i = ycols.index(col)
            return ycols[i - 1] if i >= 1 else None
        if kind == "mw":
            i = mwcols.index(col)
            return mwcols[i - 1] if i >= 1 else last_y
        if kind == "cs":
            i = cscols.index(col)
            return cscols[i - 1] if i >= 1 else last_y
        return None

    def actual(block_key, key):
        """실적 칸은 JSON 값, 나머지는 None."""
        qs = series(Q, key, len(qlab))
        ys = series(Y, key, len(ylab))

        def fn(c, kind, l):
            if kind == "q":
                return scaled(qs[qlab.index(l)], scale)
            if kind == "y":
                return scaled(ys[[str(x) for x in ylab].index(l)], scale)
            return None
        return fn

    def cs_value(key, per_share=False):
        vals = series(cs, key, len(cslab))

        def fn(col):
            v = vals[cscols.index(col)]
            return v if per_share else scaled(v, scale)
        return fn

    def yoy(key):
        def fn(c, kind, l):
            p = prev_col(c, kind)
            if p is None:
                return None
            cur, prv = S.ref(key, c), S.ref(key, p)
            return blank_safe(f"{cur}/{prv}-1", cur, prv)
        return fn

    def yoy_pos(key):  # 이익 줄 — 전년이 0 이하(적자)면 빈칸(2026-10-28 사용자: 손익 YoY 추가)
        def fn(c, kind, l):
            p = prev_col(c, kind)
            if p is None:
                return None
            cur, prv = S.ref(key, c), S.ref(key, p)
            return f'=IF(OR({cur}="",{prv}=""),"",IF({prv}<=0,"",IFERROR({cur}/{prv}-1,"")))'
        return fn

    def ratio(num, den):
        def fn(c, kind, l):
            a, b = S.ref(num, c), S.ref(den, c)
            return blank_safe(f"{a}/{b}", a, b)
        return fn

    # ── 매출 ────────────────────────────────────────────────────────────
    segs = d.get("segments") or []
    seg_keys = []
    rev_actual = actual("", "revenue")
    cs_rev = cs_value("revenue")

    def rev_fn(c, kind, l):
        if kind in ("q", "y"):
            return rev_actual(c, kind, l)
        if kind == "cs":
            return cs_rev(c)
        # MW: 제품별 매출 합계 (제품 칸이 비어 있으면 빈칸 → 매출칸에 직접 입력)
        if seg_keys:
            refs = [S.ref(k, c) for k in seg_keys]
            rng = ",".join(refs)
            return f'=IF(COUNT({rng})=0,"",SUM({rng}))'
        return None

    # 매출액 행은 제품 행 번호를 알아야 하므로 먼저 자리만 잡고 나중에 채운다
    rev_row = S.r
    S.rows["revenue"] = rev_row
    S.r += 1
    S.row("rev_yoy", "YoY", yoy("revenue"), "ratio", PCT)

    S.section("제품별 매출 분류", "사업보고서 매출실적")
    if not segs:
        segs = [{"name": "(제품별 매출 미공시 — 직접 입력)"}]
    for i, sg in enumerate(segs):
        k = f"seg{i}"
        seg_keys.append(k)
        sq = list(sg.get("quarterly") or []) + [None] * len(qlab)
        sy = list(sg.get("annual") or []) + [None] * len(ylab)
        scs = list(sg.get("consensus") or []) + [None] * len(cslab)

        def seg_fn(c, kind, l, sq=sq, sy=sy, scs=scs):
            if kind == "q":
                return scaled(sq[qlab.index(l)], scale)
            if kind == "y":
                return scaled(sy[[str(x) for x in ylab].index(l)], scale)
            if kind == "cs":
                return scaled(scs[cscols.index(c)], scale)
            return None
        S.row(k, sg["name"], seg_fn, "detail", NUM, note=sg.get("note"), input_kinds=("mw",))
        S.row(k + "_w", "비중", lambda c, kind, l, k=k: blank_safe(
            f"{S.ref(k, c)}/{S.ref('revenue', c)}", S.ref(k, c), S.ref("revenue", c)), "ratio", PCT)
        S.row(k + "_yoy", "YoY", yoy(k), "ratio", PCT)

    # 매출액 행 채우기
    S.r, save_r = rev_row, S.r
    S.row("revenue", "매출액", rev_fn, "key", NUM,
          note="MW = 제품 합계" if segs and segs[0].get("annual") else "MW 직접 입력",
          input_kinds=() if (segs and segs[0].get("annual")) else ("mw",))
    S.r = save_r

    # ── 매출원가 ────────────────────────────────────────────────────────
    S.section("매출원가")
    S.row("cogs", "전체 매출원가", actual("", "cogs"), "detail", NUM)
    S.row("cogs_yoy", "YoY", yoy_pos("cogs"), "ratio", PCT)

    # 제품별 매출원가 — 공시가 매출과 같은 분류로 원가를 나눠 줄 때만 행이 생긴다
    for i, sg in enumerate(segs):
        cq = list(sg.get("cogs_quarterly") or [])
        cy = list(sg.get("cogs_annual") or [])
        if all(v is None for v in cq + cy):
            continue
        k = f"seg{i}"
        cq += [None] * len(qlab)
        cy += [None] * len(ylab)

        def seg_cogs_fn(c, kind, l, cq=cq, cy=cy):
            if kind == "q":
                return scaled(cq[qlab.index(l)], scale)
            if kind == "y":
                return scaled(cy[[str(x) for x in ylab].index(l)], scale)
            return None
        S.row(k + "_cogs", sg["name"], seg_cogs_fn, "detail", NUM, note="제품별 매출원가")
        S.row(k + "_gpm", "GPM", lambda c, kind, l, k=k: blank_safe(
            f"1-{S.ref(k + '_cogs', c)}/{S.ref(k, c)}", S.ref(k + "_cogs", c), S.ref(k, c))
            if kind in ("q", "y") else None, "ratio", PCT)

    def gp_fn(c, kind, l):
        if kind not in ("q", "y"):
            return None
        rv, cg = S.ref("revenue", c), S.ref("cogs", c)
        return blank_safe(f"{rv}-{cg}", rv, cg)
    S.row("gp", "매출총이익", gp_fn, "plain", NUM)
    S.row("gpm", "GPM", ratio("gp", "revenue"), "ratio", PCT)
    S.row("gp_yoy", "YoY", yoy_pos("gp"), "ratio", PCT)

    # ── 판관비·영업이익 ──────────────────────────────────────────────────
    S.section("판매관리비", sub=True)
    S.row("sga", "전체 판관비", actual("", "sga"), "detail", NUM)
    S.row("sga_yoy", "YoY", yoy_pos("sga"), "ratio", PCT)
    op_actual = actual("", "op")
    cs_op = cs_value("op")
    S.rows["opm"] = S.r + 1  # 영업이익 수식이 OPM 행을 참조

    def op_fn(c, kind, l):
        if kind in ("q", "y"):
            return op_actual(c, kind, l)
        if kind == "cs":
            return cs_op(c)
        rv, om = S.ref("revenue", c), f"{CL(c)}{S.rows['opm']}"
        return blank_safe(f"{rv}*{om}", rv, om)
    S.row("op", "영업이익", op_fn, "key", NUM, note="MW = 매출 × OPM")

    def opm_fn(c, kind, l):
        if kind == "mw":
            return None  # 입력칸
        a, b = S.ref("op", c), S.ref("revenue", c)
        return blank_safe(f"{a}/{b}", a, b)
    S.row("opm", "OPM", opm_fn, "ratio", PCT, note="MW 입력", input_kinds=("mw",))
    S.row("op_yoy", "증가율", yoy("op"), "ratio", PCT)

    # ── 비용의 성격별 분류 (연간만) ──────────────────────────────────────
    S.section("비용의 성격별 분류", "참고 데이터")

    def total_sga(c, kind, l):
        return (f'=IF({S.ref("sga", c)}="","",{S.ref("sga", c)})' if kind == "y" else None)
    S.row("sga_tot", "총 판관비", total_sga, "key", NUM)
    nature = [("salary", "급여", "고정비"), ("dep", "감가상각비", "고정비"),
              ("fees", "지급수수료", "고정비"), ("rnd", "연구개발비", "변동비")]
    ys_all = {k: series(Y, k, len(ylab)) for k, _, _ in nature}
    for k, name, kind_note in nature:
        vals = ys_all[k]

        def nat_fn(c, kind, l, vals=vals):
            if kind != "y":
                return None
            return scaled(vals[[str(x) for x in ylab].index(l)], scale)
        S.row(k, name, nat_fn, "detail", NUM, note=kind_note)
        S.row(k + "_pct", "% of Sales", lambda c, kind, l, k=k: blank_safe(
            f"{S.ref(k, c)}/{S.ref('revenue', c)}", S.ref(k, c), S.ref("revenue", c))
            if kind == "y" else None, "ratio", PCT)

    def other_cost(c, kind, l):
        if kind != "y":
            return None
        parts = [S.ref(k, c) for k, _, _ in nature]
        tot = S.ref("sga_tot", c)
        return blank_safe(f"{tot}-" + "-".join(parts), tot, *parts)
    S.row("other_cost", "기타비용", other_cost, "detail", NUM, note="변동비")
    S.row("other_cost_pct", "% of Sales", lambda c, kind, l: blank_safe(
        f"{S.ref('other_cost', c)}/{S.ref('revenue', c)}", S.ref("other_cost", c), S.ref("revenue", c))
        if kind == "y" else None, "ratio", PCT)

    # ── 당기순이익 추정 ──────────────────────────────────────────────────
    S.section("당기순이익 추정")
    recent = ycols[-3:] if len(ycols) >= 3 else ycols
    for k, name in [("oth_inc", "기타이익"), ("oth_exp", "기타손실"),
                    ("fin_inc", "금융수익"), ("fin_exp", "금융원가")]:
        name = meta.get("labels", {}).get(k, name)
        act = actual("", k)
        pct_row = S.r + 1

        def item_fn(c, kind, l, act=act, pct_row=pct_row):
            if kind in ("q", "y"):
                return act(c, kind, l)
            if kind == "mw":
                rv, p = S.ref("revenue", c), f"{CL(c)}{pct_row}"
                return blank_safe(f"{rv}*{p}", rv, p)
            return None
        S.row(k, name, item_fn, "detail", NUM)

        def pct_fn(c, kind, l, k=k):
            if kind in ("q", "y"):
                a, b = S.ref(k, c), S.ref("revenue", c)
                return blank_safe(f"{a}/{b}", a, b)
            if kind == "mw":
                i = mwcols.index(c)
                if i == 0:
                    rng = ",".join(f"{CL(x)}{S.r}" for x in recent)
                    return f'=IFERROR(AVERAGE({rng}),"")'
                return f"={CL(mwcols[i - 1])}{S.r}"
            return None
        S.row(k + "_pct", "% of Sales", pct_fn, "ratio", PCT,
              note="MW = 최근 3년 평균" if k == "oth_inc" else None, input_kinds=("mw",))

    pbt_actual = actual("", "pbt")
    cs_pbt = cs_value("pbt")

    def pbt_fn(c, kind, l):
        if kind in ("q", "y"):
            return pbt_actual(c, kind, l)
        if kind == "cs":
            return cs_pbt(c)
        refs = [S.ref(k, c) for k in ("op", "oth_inc", "oth_exp", "fin_inc", "fin_exp")]
        o, oi, oe, fi, fe = refs
        return blank_safe(f"{o}+N({oi})-N({oe})+N({fi})-N({fe})", o)
    S.row("pbt", "법인세비용차감전순이익(손실)", pbt_fn, "key", NUM)
    S.row("pbt_yoy", "YoY", yoy_pos("pbt"), "ratio", PCT)
    tax_actual = actual("", "tax")
    S.rows["tax_rate"] = S.r + 1

    def tax_fn(c, kind, l):
        if kind in ("q", "y"):
            return tax_actual(c, kind, l)
        if kind == "mw":
            p, t = S.ref("pbt", c), f"{CL(c)}{S.rows['tax_rate']}"
            return blank_safe(f"MAX({p},0)*{t}", p, t)
        return None
    S.row("tax", "법인세비용(수익)", tax_fn, "detail", NUM)
    default_tax = meta.get("tax_rate", DEFAULT_TAX.get(mkt, 0.22))

    def tax_rate_fn(c, kind, l):
        if kind in ("q", "y"):
            a, b = S.ref("tax", c), S.ref("pbt", c)
            return blank_safe(f"{a}/{b}", a, b)
        if kind == "mw":
            return default_tax
        return None
    S.row("tax_rate", "법인세율", tax_rate_fn, "ratio", PCT,
          note=f"MW 기본 {default_tax:.0%}", input_kinds=("mw",))

    ni_actual = actual("", "ni")
    cs_ni = cs_value("ni")

    def ni_fn(c, kind, l):
        if kind in ("q", "y"):
            return ni_actual(c, kind, l)
        if kind == "cs":
            return cs_ni(c)
        p, t = S.ref("pbt", c), S.ref("tax", c)
        return blank_safe(f"{p}-{t}", p)
    S.row("ni", "당기순이익(손실)", ni_fn, "key", NUM)
    S.row("ni_yoy", "YoY", yoy_pos("ni"), "ratio", PCT)
    own_actual = actual("", "ni_owner")
    cs_own = cs_value("ni_owner")
    S.rows["own_ratio"] = S.r + 1

    def own_fn(c, kind, l):
        if kind in ("q", "y"):
            return own_actual(c, kind, l)
        if kind == "cs":
            return cs_own(c)
        n, rt = S.ref("ni", c), f"{CL(c)}{S.rows['own_ratio']}"
        return blank_safe(f"{n}*{rt}", n, rt)
    S.row("ni_owner", "지배주주 순이익", own_fn, "key", NUM)

    def own_ratio_fn(c, kind, l):
        if kind == "mw":
            return 1
        a, b = S.ref("ni_owner", c), S.ref("ni", c)
        return blank_safe(f"{a}/{b}", a, b)
    S.row("own_ratio", "지배주주 순이익률", own_ratio_fn, "ratio", PCT,
          note="MW 기본 100%", input_kinds=("mw",))
    S.row("ni_owner_yoy", "지배주주 순이익 YoY", yoy_pos("ni_owner"), "ratio", PCT)

    # ── 목표주가 산출 ────────────────────────────────────────────────────
    mk = d.get("market_data") or {}
    S.section("목표주가 산출")
    price_r = S.row("price", f"현재 주가 ({mk.get('price_date', '')})", lambda *a: None, "detail", WON)
    ws.cell(price_r, 2, mk.get("price")).number_format = WON
    shares_r = S.row("shares", "발행주식수 (보통주)", lambda *a: None, "detail", WON)
    ws.cell(shares_r, 2, mk.get("shares")).number_format = WON
    mcap_r = S.row("mcap", f"시가총액 ({meta['unit_label']})", lambda *a: None, "detail", NUM)
    ws.cell(mcap_r, 2, f'=IF(OR(B{price_r}="",B{shares_r}=""),"",B{price_r}*B{shares_r}/{scale})').number_format = NUM
    for r_ in (price_r, shares_r, mcap_r):
        ws.cell(r_, 2).font = font(True)

    S.row("v_opm", "OPM", lambda c, kind, l: f'=IF({S.ref("opm", c)}="","",{S.ref("opm", c)})'
          if kind in ("y", "mw", "cs") else None, "ratio", PCT)
    S.row("v_own_yoy", "지배주주 순이익 증가율", lambda c, kind, l: yoy("ni_owner")(c, kind, l)
          if kind in ("y", "mw", "cs") else None, "ratio", PCT)

    def pe_fn(c, kind, l):
        if kind == "y" and c != last_y:
            return None
        if kind == "q":
            return None
        n = S.ref("ni_owner", c)
        return blank_safe(f"$B${mcap_r}/{n}", f"$B${mcap_r}", n)
    S.row("pe", "P/E (현재 주가)", pe_fn, "plain", MULT, note="시총 ÷ 지배순이익")
    S.row("tpe", "Target P/E", lambda *a: None, "plain", MULT, note="직접 입력", input_kinds=("mw", "cs"))
    cs_eps = cs_value("eps", per_share=True)

    def eps_fn(c, kind, l):
        if kind == "q":
            return None
        if kind == "cs" and cs_eps(c) is not None:
            return cs_eps(c)
        n = S.ref("ni_owner", c)
        return blank_safe(f"{n}*{scale}/$B${shares_r}", n, f"$B${shares_r}")
    S.row("eps", "EPS", eps_fn, "plain", WON, note="지배순이익 ÷ 주식수")

    def tp_fn(c, kind, l):
        if kind not in ("mw", "cs"):
            return None
        e, t = S.ref("eps", c), S.ref("tpe", c)
        return blank_safe(f"{e}*{t}", e, t)
    S.row("tp", "목표주가", tp_fn, "key", WON)

    def up_fn(c, kind, l):
        if kind not in ("mw", "cs"):
            return None
        t = S.ref("tp", c)
        return blank_safe(f"{t}/$B${price_r}-1", t, f"$B${price_r}")
    r_up = S.row("upside", "상승여력", up_fn, "plain", PCT)
    for c in mwcols + cscols:
        ws.cell(r_up, c).font = font(True, RED)

    # ── 자산가치·기업가치 (재무상태표 시트 최신 분기 참조) ────────────────
    S.section("자산가치 산출", "최신 분기 BS")
    bs_latest = d.get("_bs_latest_ref", {})
    tail = ([last_y] if last_y else []) + mwcols

    def bs_link(key):
        def fn(c, kind, l):
            if c not in tail or key not in bs_latest:
                return None
            return f"={bs_latest[key]}"
        return fn
    S.row("a_cash", "현금및현금성자산", bs_link("cash"), "detail", NUM)
    S.row("a_stfin", "단기금융상품", bs_link("st_fin"), "detail", NUM)
    S.row("a_prop", "투자부동산", bs_link("inv_property"), "detail", NUM)
    S.row("a_eq", "지배자본", bs_link("equity_owner"), "detail", NUM)
    S.row("a_liab", "부채", bs_link("total_liab"), "detail", NUM)
    S.row("a_dr", "부채비율", lambda c, kind, l: blank_safe(
        f"{S.ref('a_liab', c)}/{S.ref('a_eq', c)}", S.ref("a_liab", c), S.ref("a_eq", c))
        if c in tail else None, "ratio", PCT)
    S.row("a_val", "자산가치", lambda c, kind, l: f"=SUM({S.ref('a_cash', c)},{S.ref('a_stfin', c)},{S.ref('a_prop', c)})"
          if c in tail else None, "key", NUM, note="현금+단기금융+투자부동산")

    S.section("기업가치 산출")

    def tmcap_fn(c, kind, l):
        if kind not in ("mw", "cs"):
            return None
        n, t = S.ref("ni_owner", c), S.ref("tpe", c)
        return blank_safe(f"{n}*{t}", n, t)
    S.row("t_mcap", "목표 시가총액 (P/E)", tmcap_fn, "plain", NUM, note="지배순이익 × Target P/E")

    def ev_fn(c, kind, l):
        if kind != "mw":
            return None
        t, a = S.ref("t_mcap", c), S.ref("a_val", c)
        return blank_safe(f"{t}+{a}", t)
    S.row("ev", "기업가치", ev_fn, "key", NUM, note="목표시총 + 자산가치")
    S.row("cur_mcap", "현재 시가총액", lambda c, kind, l: f"=$B${mcap_r}" if kind == "mw" else None, "plain", NUM)
    r_up2 = S.row("ev_up", "상승여력", lambda c, kind, l: blank_safe(
        f"{S.ref('ev', c)}/{S.ref('cur_mcap', c)}-1", S.ref("ev", c), S.ref("cur_mcap", c))
        if kind == "mw" else None, "plain", PCT)
    for c in mwcols:
        ws.cell(r_up2, c).font = font(True, RED)

    # 출처
    S.gap()
    for line in ["출처"] + list(meta.get("sources", [])):
        ws.cell(S.r, 1, line).font = font(line == "출처", GRAY_TXT, 9)
        ws.cell(S.r, 1).alignment = Alignment(horizontal="left")
        S.r += 1

    # 블록 경계선 (분기 | 연간 | MW | 컨센서스)
    for block in (qcols, ycols, mwcols, cscols):
        if block:
            for r in range(3, S.r):
                cell = ws.cell(r, block[0])
                b = cell.border
                cell.border = Border(left=thin, top=b.top, bottom=b.bottom)
    for c in mwcols:
        ws.cell(3, c).fill = fill("DDEBF7")
    for c in cscols:
        ws.cell(3, c).fill = fill("FFF2CC")

    group_cols(ws, qcols)
    ws.freeze_panes = "C4"
    print_setup(ws)
    return S, qcol, ycol


# ── 시트 2: 재무상태표 + 주요 재무 비율 ───────────────────────────────────
BS_ROWS = [
    ("section", "자산"),
    ("current_assets", "유동자산", "key"),
    ("cash", "현금및현금성자산", "detail"),
    ("st_fin", "단기금융상품", "detail"),
    ("receivables", "매출채권", "detail"),
    ("inventory", "재고자산", "detail"),
    ("other_ca", "기타유동자산", "detail", "calc"),
    ("noncurrent_assets", "비유동자산", "key"),
    ("ppe", "유형자산", "detail"),
    ("intangibles", "무형자산", "detail"),
    ("inv_property", "투자부동산", "detail"),
    ("other_nca", "기타비유동자산", "detail", "calc"),
    ("total_assets", "자산총계", "key"),
    ("section", "부채"),
    ("current_liab", "유동부채", "key"),
    ("payables", "매입채무", "detail"),
    ("st_debt", "단기차입금", "detail"),
    ("current_ltd", "유동성장기부채", "detail"),
    ("other_cl", "기타유동부채", "detail", "calc"),
    ("noncurrent_liab", "비유동부채", "key"),
    ("lt_debt", "장기차입금", "detail"),
    ("bonds", "사채", "detail"),
    ("other_ncl", "기타비유동부채", "detail", "calc"),
    ("total_liab", "부채총계", "key"),
    ("total_debt", "총차입금", "plain", "calc"),
    ("net_debt", "순차입금", "plain", "calc"),
    ("section", "자본"),
    ("equity_owner", "지배주주지분", "key"),
    ("capital_stock", "자본금", "detail"),
    ("retained", "이익잉여금", "detail"),
    ("nci", "비지배지분", "detail"),
    ("total_equity", "자본총계", "key"),
]

CALC = {
    "other_ca": ("current_assets", ["cash", "st_fin", "receivables", "inventory"]),
    "other_nca": ("noncurrent_assets", ["ppe", "intangibles", "inv_property"]),
    "other_cl": ("current_liab", ["payables", "st_debt", "current_ltd"]),
    "other_ncl": ("noncurrent_liab", ["lt_debt", "bonds"]),
}

# (key, 라벨, 기준(참고칸), 서식, 수식 종류, 불합격 조건)
RATIOS = [
    ("sub", "안정성"),
    ("r_debt", "부채비율", "100%↓", PCT, ("bs", "total_liab", "total_equity"), (">", 1.0)),
    ("r_netdebt", "순차입금비율", "0%↓ 순현금", PCT, ("bs", "net_debt", "total_equity"), (">", 0.5)),
    ("r_eqratio", "자기자본비율", "50%↑", PCT, ("bs", "total_equity", "total_assets"), ("<", 0.5)),
    ("r_current", "유동비율", "150%↑", PCT, ("bs", "current_assets", "current_liab"), ("<", 1.5)),
    ("r_quick", "당좌비율", "100%↑", PCT, ("quick",), ("<", 1.0)),
    ("r_icr", "이자보상배율", "5x↑ (금융원가 기준)", MULT, ("flow_flow", "op", "fin_exp"), ("<", 5)),
    ("r_payback", "차입금 상환연수", "4년↓ (버핏)", YRS, ("bs_flow", "total_debt", "ni_owner"), (">", 4)),
    ("sub", "수익성"),
    ("r_roe", "ROE", "15%↑ (버핏)", PCT, ("flow_bs", "ni_owner", "equity_owner"), ("<", 0.15)),
    ("r_roa", "ROA", "5%↑", PCT, ("flow_bs", "ni", "total_assets"), ("<", 0.05)),
    ("r_roic", "ROIC", "10%↑ (세후영업이익 기준)", PCT, ("roic",), ("<", 0.10)),
    ("r_gpm", "매출총이익률 (GPM)", "40%↑ (해자 신호)", PCT, ("flow_flow", "gp", "revenue"), ("<", 0.40)),
    ("r_opm", "영업이익률 (OPM)", "15%↑", PCT, ("flow_flow", "op", "revenue"), ("<", 0.15)),
    ("r_npm", "순이익률", "10%↑", PCT, ("flow_flow", "ni", "revenue"), ("<", 0.10)),
    ("sub", "효율성 (회전율)"),
    ("r_ato", "총자산회전율", "높을수록↑", MULT, ("flow_bs", "revenue", "total_assets"), None),
    ("r_ar_days", "매출채권회전일수", "짧을수록↑", DAYS, ("days", "receivables", "revenue"), None),
    ("r_inv_days", "재고자산회전일수", "짧을수록↑", DAYS, ("days", "inventory", "cogs"), None),
    ("r_ap_days", "매입채무회전일수", "길수록 유리", DAYS, ("days", "payables", "cogs"), None),
    ("r_ccc", "현금전환주기 (CCC)", "짧을수록↑", DAYS, ("ccc",), None),
    ("sub", "현금흐름"),
    ("r_ocf_ni", "영업현금흐름 / 순이익", "1x↑ (이익의 질)", MULT, ("cf_flow", "ocf", "ni"), ("<", 1.0)),
    ("r_fcf_m", "FCF 마진", "5%↑", PCT, ("cf_flow", "fcf", "revenue"), ("<", 0.05)),
    ("r_capex", "CAPEX / 매출", "낮을수록↑ (자본 가벼움)", PCT, ("cf_flow", "capex", "revenue"), None),
    ("r_payout", "배당성향", "참고", PCT, ("cf_flow", "dividends", "ni_owner"), None),
]


def build_bs(wb, d, A, qcol_a, ycol_a):
    meta, scale = d["meta"], d["meta"]["scale"]
    mkt = meta.get("market", "KR")
    bs = d.get("bs", {})
    BQ, BY = bs.get("quarterly", {}), bs.get("annual", {})
    bq_lab, by_lab = BQ.get("labels", []), [str(x) for x in BY.get("labels", [])]
    cf = d.get("cf", {})
    CFY, CFQ = cf.get("annual", {}), cf.get("quarterly", {})

    ws = wb.create_sheet("재무상태표")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    cols, c = [], 3
    for l in bq_lab:
        cols.append((c, "q", l)); c += 1
    for l in by_lab:
        cols.append((c, "y", l)); c += 1
    S = Sheet(ws, cols, label_width=24)
    ws.column_dimensions["B"].width = 20

    ws.cell(1, 1, f"{meta['company']} 재무상태표").font = font(True, size=20)
    ws.row_dimensions[1].height = 30
    ws.cell(2, 1, "재무비율은 기말 잔액 기준 · 분기 손익은 최근 4분기 합산(TTM)").font = font(False, RED, 9)
    ws.cell(2, 2, f"(단위: {meta['unit_label']})").font = font(False, GRAY_TXT, 9)
    ws.cell(3, 1, "항목"); ws.cell(3, 2, "참고")
    for cc, k, l in cols:
        ws.cell(3, cc, l)
    for cc in range(1, S.last_col + 1):
        cell = ws.cell(3, cc)
        cell.fill = fill(GRAY_HDR)
        cell.font = font(True, size=10)
        cell.alignment = CENTER
        cell.border = Border(top=thin, bottom=thin)

    def bs_val(key):
        qv, yv = series(BQ, key, len(bq_lab)), series(BY, key, len(by_lab))

        def fn(cc, kind, l):
            if kind == "q":
                return scaled(qv[bq_lab.index(l)], scale)
            return scaled(yv[by_lab.index(l)], scale)
        return fn

    for spec in BS_ROWS:
        if spec[0] == "section":
            S.section(spec[1]); continue
        key, label, style = spec[0], spec[1], spec[2]
        is_calc = len(spec) > 3
        if is_calc and key in CALC:
            tot, parts = CALC[key]

            def fn(cc, kind, l, tot=tot, parts=parts):
                t = S.ref(tot, cc)
                ps = [S.ref(p, cc) for p in parts]
                return blank_safe(f"{t}-SUM({','.join(ps)})", t)
            S.row(key, label, fn, style, NUM, note="총계 − 위 항목")
        elif key == "total_debt":
            S.row(key, label, lambda cc, kind, l: f"=SUM({','.join(S.ref(p, cc) for p in ('st_debt', 'current_ltd', 'lt_debt', 'bonds'))})",
                  style, NUM, note="단기+유동성장기+장기+사채")
        elif key == "net_debt":
            S.row(key, label, lambda cc, kind, l: f"={S.ref('total_debt', cc)}-{S.ref('cash', cc)}-N({S.ref('st_fin', cc)})",
                  style, NUM, note="음수 = 순현금")
        else:
            S.row(key, label, bs_val(key), style, NUM)

    # 현금흐름 요약
    S.section("현금흐름 요약", "분기 = TTM (있을 때만)")

    def cf_val(key):
        qv, yv = series(CFQ, key, len(bq_lab)), series(CFY, key, len(by_lab))
        ylabs = [str(x) for x in CFY.get("labels", by_lab)]

        def fn(cc, kind, l):
            if kind == "q":
                return scaled(qv[bq_lab.index(l)], scale) if CFQ else None
            if l in ylabs:
                return scaled(series(CFY, key, len(ylabs))[ylabs.index(l)], scale)
            return None
        return fn
    S.row("ocf", "영업활동현금흐름", cf_val("ocf"), "key", NUM)
    S.row("capex", "CAPEX (유·무형 취득)", cf_val("capex"), "detail", NUM, note="양수로 입력")
    S.row("fcf", "FCF", lambda cc, kind, l: blank_safe(f"{S.ref('ocf', cc)}-{S.ref('capex', cc)}",
                                                        S.ref("ocf", cc), S.ref("capex", cc)), "key", NUM)
    S.row("dividends", "배당금 지급", cf_val("dividends"), "detail", NUM, note="양수로 입력")

    # 연간 시트 손익 참조
    def flow(key, cc, kind, l):
        """손익 항목: 연간은 해당 연도, 분기는 TTM(4분기 합)."""
        if key not in A.rows:
            return None
        r = A.rows[key]
        if kind == "y":
            col = ycol_a.get(l)
            return f"'연간'!{CL(col)}{r}" if col else None
        col = qcol_a.get(l)
        if not col:
            return None
        qcols = sorted(qcol_a.values())
        i = qcols.index(col)
        if i < 3:
            return None
        rng = f"'연간'!{CL(qcols[i - 3])}{r}:{CL(col)}{r}"
        return f'IF(COUNT({rng})<4,"",SUM({rng}))'

    def cf_flow(key, cc):
        return S.ref(key, cc)

    S.section("주요 재무 비율", "기준 (빨간 글씨 = 미달)")
    for spec in RATIOS:
        if spec[0] == "sub":
            S.section(spec[1], sub=True); continue
        key, label, bench, fmt, kind_, fail = spec

        def fn(cc, kind, l, kind_=kind_):
            t = kind_[0]
            try:
                if t == "bs":
                    a, b = S.ref(kind_[1], cc), S.ref(kind_[2], cc)
                    return blank_safe(f"{a}/{b}", a, b)
                if t == "quick":
                    ca, inv, cl = S.ref("current_assets", cc), S.ref("inventory", cc), S.ref("current_liab", cc)
                    return blank_safe(f"({ca}-N({inv}))/{cl}", ca, cl)
                if t == "flow_flow":
                    a, b = flow(kind_[1], cc, kind, l), flow(kind_[2], cc, kind, l)
                    if not a or not b:
                        return None
                    return f'=IFERROR(({a})/({b}),"")'
                if t == "flow_bs":
                    a, b = flow(kind_[1], cc, kind, l), S.ref(kind_[2], cc)
                    if not a:
                        return None
                    return f'=IFERROR(({a})/{b},"")'
                if t == "bs_flow":
                    a, b = S.ref(kind_[1], cc), flow(kind_[2], cc, kind, l)
                    if not b:
                        return None
                    return f'=IFERROR(IF(({b})<=0,"적자",{a}/({b})),"")'
                if t == "days":
                    a, b = S.ref(kind_[1], cc), flow(kind_[2], cc, kind, l)
                    if not b:
                        return None
                    return f'=IFERROR({a}/({b})*365,"")'
                if t == "ccc":
                    ar, inv, ap = (S.ref(k, cc) for k in ("r_ar_days", "r_inv_days", "r_ap_days"))
                    return blank_safe(f"{ar}+{inv}-{ap}", ar, inv, ap)
                if t == "roic":
                    op = flow("op", cc, kind, l)
                    if not op:
                        return None
                    te, nd = S.ref("total_equity", cc), S.ref("net_debt", cc)
                    tax = meta.get("tax_rate", DEFAULT_TAX.get(mkt, 0.22))
                    return f'=IFERROR(({op})*(1-{tax})/({te}+{nd}),"")'
                if t == "cf_flow":
                    a = cf_flow(kind_[1], cc)
                    b = flow(kind_[2], cc, kind, l)
                    if not b:
                        return None
                    return f'=IF({a}="","",IFERROR({a}/({b}),""))'
            except KeyError:
                return None
            return None
        r = S.row(key, label, fn, "plain", fmt, note=bench)
        ws.cell(r, 1).alignment = LEFT
        if fail:
            op_, th = fail
            first, last = CL(cols[0][0]), CL(cols[-1][0])
            rng = f"{first}{r}:{last}{r}"
            ws.conditional_formatting.add(
                rng, FormulaRule(formula=[f"AND(ISNUMBER({first}{r}),{first}{r}{op_}{th})"],
                                 font=Font(name=FONT, color=RED, bold=True)))

    for block in ([c for c, k, _ in cols if k == "q"], [c for c, k, _ in cols if k == "y"]):
        if block:
            for r in range(3, S.r):
                cell = ws.cell(r, block[0])
                b = cell.border
                cell.border = Border(left=thin, top=b.top, bottom=b.bottom)
    group_cols(ws, [c for c, k, _ in cols if k == "q"])
    ws.freeze_panes = "C4"
    print_setup(ws)

    d["_bs_sheet"] = (S, cols)
    # 연간 시트 자산가치가 참조할 최신 분기 칸
    latest_col = cols[0][0] + len(bq_lab) - 1 if bq_lab else (cols[-1][0] if cols else None)
    refs = {}
    if latest_col:
        for k in ("cash", "st_fin", "inv_property", "equity_owner", "total_liab"):
            refs[k] = f"'재무상태표'!{CL(latest_col)}{S.rows[k]}"
    return refs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json_path")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args()
    with open(args.json_path, encoding="utf-8") as f:
        d = json.load(f)

    # 재무상태표 참조 주소를 먼저 알아야 연간 시트 자산가치 수식을 쓸 수 있다 → 2패스
    wb = Workbook()
    A, qcol_a, ycol_a = build_annual(wb, d)
    refs = build_bs(wb, d, A, qcol_a, ycol_a)
    d["_bs_latest_ref"] = refs
    wb = Workbook()
    A, qcol_a, ycol_a = build_annual(wb, d)
    build_bs(wb, d, A, qcol_a, ycol_a)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.out)
    print(f"saved: {args.out}")


if __name__ == "__main__":
    sys.exit(main())
